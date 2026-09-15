"""Generate five synthetic contracts that exercise content-first applicability.

WHY THIS EXISTS. `AM-51`/`AM-60` made applicability a question about what a
document CONTAINS, not what its label says. Nothing in the corpus proved that on
a realistic document: the STRUCTURAL fixture is one short synthetic clause set,
and the real counterparty paper is MSA/TOS/NDA-shaped and gitignored.

These five are **test inputs, never legal positions** (rule 21). Every number in
them is chosen to land on a KNOWN side of a ratified Company Standard so the
expected outcome is derived from the standard, not asserted by me:

    C1  MSA-labelled,   every clause on the Constitution's number   -> all Acceptable
    C2  OTHER-labelled, the SAME clauses                            -> all Acceptable
                        (the headline: the label must not change the answer)
    C3  MSA-labelled,   some on the number, some off it             -> mixed
    C4  OTHER-labelled, MSA + NDA clauses mixed in one document     -> mixed
    C5  OTHER-labelled, an UNCAPPED liability term                  -> Needs a decision
                        (`constitution_boundaries` §9: the one prohibition the
                         liability evaluator already extracts)

Synthetic or cleared text only — CI job 8 rejects a real contract entering the
repository, and locked 54.6 keeps documents out of it entirely. The `.docx`
files are written into `legal-docs/` (gitignored); only this generator is
committed, exactly as `e2e_bootstrap.py` does for its STRUCTURAL fixture.

No counterparty is real. "Northwind Systems Private Limited" and the rest are
invented names; any resemblance to a real company is accidental and unintended.

Usage:
    python3 -m tools.generate_test_contracts --out ../legal-docs/synthetic
"""

from __future__ import annotations

import argparse
from pathlib import Path

#: Boilerplate that carries no position — it exists so each document reads like
#: a real agreement and runs to four or five pages, which is what makes the
#: segmentation and mapping work non-trivial.
FILLER = {
    "definitions": [
        '"Affiliate" means any entity that directly or indirectly controls, is '
        "controlled by, or is under common control with a party, where "
        '"control" means ownership of more than fifty percent (50%) of the '
        "voting securities of that entity.",
        '"Confidential Information" means all non-public information disclosed '
        "by one party to the other, whether orally, visually, in writing or in "
        "any other tangible or intangible form, that is designated as "
        "confidential at the time of disclosure or that a reasonable person "
        "would understand to be confidential given the nature of the "
        "information and the circumstances of its disclosure.",
        '"Deliverables" means the reports, configurations, documentation and '
        "other work product prepared by the Supplier specifically for the "
        "Customer under a Statement of Work.",
        '"Services" means the services described in the applicable Statement '
        "of Work, together with any support, maintenance and professional "
        "services the parties agree in writing to include.",
        '"Statement of Work" or "SOW" means a document executed by both parties '
        "that describes the Services, the fees, the timetable and any "
        "acceptance criteria.",
    ],
    "scope": [
        "The Supplier shall perform the Services with reasonable skill and care "
        "and in accordance with the standards generally observed in the "
        "industry for similar services. The Supplier shall allocate personnel "
        "with the qualifications and experience reasonably necessary to perform "
        "the Services.",
        "Each Statement of Work forms part of this Agreement and is subject to "
        "its terms. In the event of a conflict between the body of this "
        "Agreement and a Statement of Work, the body of this Agreement prevails "
        "unless the Statement of Work expressly states otherwise and is signed "
        "by an authorised representative of each party.",
        "The Customer shall provide the Supplier with such access, information, "
        "facilities and cooperation as the Supplier reasonably requires to "
        "perform the Services. The Supplier is not liable for a delay or "
        "failure to perform to the extent caused by the Customer's failure to "
        "provide that cooperation.",
    ],
    "warranties": [
        "Each party represents and warrants that it has full corporate power "
        "and authority to enter into this Agreement and to perform its "
        "obligations under it, and that this Agreement has been duly authorised "
        "by all necessary corporate action.",
        "Each party represents and warrants that its performance of this "
        "Agreement will not violate any agreement to which it is a party or by "
        "which it is bound, and that it holds all licences, permits and "
        "registrations required to perform its obligations.",
    ],
    "compliance": [
        "Each party shall comply with all applicable laws, rules and "
        "regulations in the performance of this Agreement, including those "
        "relating to anti-bribery and anti-corruption, economic sanctions, "
        "export control, and the protection of personal data.",
        "The Supplier shall maintain reasonable technical and organisational "
        "measures designed to protect the Customer's data against unauthorised "
        "access, loss, alteration or disclosure, and shall notify the Customer "
        "without undue delay after becoming aware of a personal data breach "
        "affecting the Customer's data.",
    ],
    "force_majeure": [
        "Neither party is liable for a failure or delay in performing its "
        "obligations under this Agreement to the extent that the failure or "
        "delay results from an event beyond its reasonable control, including "
        "an act of God, flood, fire, earthquake, epidemic, war, terrorism, "
        "riot, labour dispute, or failure of a public telecommunications "
        "network, provided the affected party notifies the other promptly and "
        "uses reasonable efforts to mitigate the effect of the event.",
    ],
    "notices": [
        "Any notice given under this Agreement shall be in writing and shall be "
        "delivered by hand, by registered post, or by electronic mail to the "
        "address set out in the preamble or to such other address as a party "
        "notifies in writing. A notice sent by electronic mail is deemed "
        "received on the next business day after transmission.",
        "The parties are independent contractors. Nothing in this Agreement "
        "creates a partnership, joint venture, agency or employment "
        "relationship between them, and neither party has authority to bind the "
        "other.",
    ],
    "general": [
        "This Agreement constitutes the entire agreement between the parties "
        "with respect to its subject matter and supersedes all prior "
        "negotiations, understandings and agreements, whether oral or written, "
        "relating to that subject matter.",
        "No amendment or waiver of any provision of this Agreement is effective "
        "unless made in writing and signed by an authorised representative of "
        "each party. A waiver of a breach is not a waiver of any subsequent "
        "breach.",
        "If any provision of this Agreement is held to be invalid or "
        "unenforceable, that provision shall be severed and the remaining "
        "provisions shall continue in full force and effect.",
        "Neither party may assign this Agreement without the prior written "
        "consent of the other, which shall not be unreasonably withheld, except "
        "that either party may assign to an Affiliate or to a successor in "
        "connection with a merger or a sale of substantially all of its assets.",
    ],
}


# ---------------------------------------------------------------------------
# The clauses that carry a position. Each takes the number as an argument, so a
# MATCH and a DEVIATION are the SAME sentence with a different figure — which
# is what makes the comparison the thing under test rather than the wording.
# ---------------------------------------------------------------------------
def liability_capped(months: int) -> tuple[str, list[str]]:
    return ("Limitation of Liability", [
        "Subject to the following paragraph, the total aggregate liability of "
        "either party arising out of or in connection with this Agreement, "
        f"whether in contract, tort (including negligence) or otherwise, shall "
        f"not exceed the total fees paid by the Customer under this Agreement "
        f"in the {months} ({months}) months immediately preceding the event "
        "giving rise to the claim.",
        "In no event shall either party be liable for any indirect, "
        "incidental, special, consequential, exemplary, or punitive damages, "
        "or for any loss of profits, loss of revenue, loss of business, or "
        "loss of anticipated savings, whether or not that party was advised of "
        "the possibility of such damages.",
        "Nothing in this Agreement excludes or limits the liability of either "
        "party for death or personal injury caused by its negligence, for "
        "fraud or fraudulent misrepresentation, or for any other liability "
        "that cannot lawfully be excluded or limited.",
    ])


def liability_uncapped() -> tuple[str, list[str]]:
    """The one condition `constitution_boundaries` §9 names as Not Negotiable."""
    return ("Limitation of Liability", [
        "Each party accepts unlimited liability for any loss or damage arising "
        "out of or in connection with this Agreement. The liability of the "
        "Supplier under this Agreement is not capped and shall not be subject "
        "to any financial limit, whether by reference to fees paid or "
        "otherwise.",
        "In no event shall either party be liable for any indirect, "
        "incidental, special, consequential, exemplary, or punitive damages, "
        "or for any loss of profits, loss of revenue, loss of business, or "
        "loss of anticipated savings.",
        "Nothing in this Agreement excludes or limits the liability of either "
        "party for death or personal injury caused by its negligence or for "
        "fraud.",
    ])


def payment(days: int, dispute_days: int = 0) -> tuple[str, list[str]]:
    del dispute_days  # see the note below: deliberately not drafted
    return ("Fees, Invoicing and Payment", [
        "The Customer shall pay the fees set out in the applicable Statement "
        "of Work. All fees are stated exclusive of GST and any other "
        "applicable taxes, levies or duties, which shall be charged in "
        "addition at the rate in force at the tax point and shall be payable "
        "by the Customer.",
        f"The Customer shall pay each undisputed invoice within {days} "
        f"({days}) days of the invoice date, in the currency stated on the "
        "invoice and without set-off, deduction or counterclaim.",
        # NO billing-dispute sentence here, deliberately.
        #
        # PAYMENT-PERIOD-MSA-001 claims cap_phrases "days of the invoice date",
        # "days of invoice" AND "days of receipt of the invoice", with the single
        # basis term "invoice". So ANY "within N days ... invoice" sentence in
        # the document is read as an INVOICE_PAYMENT_PERIOD. A billing-dispute
        # window is ordinary drafting and is indistinguishable to it: the
        # evaluator then holds two values for one basis and classifies CONFLICT,
        # which surfaces as "Needs a decision" on a payment term that is in fact
        # exactly on the Constitution's number.
        #
        # That is fail-closed and therefore SAFE (rule 15), but it is a real
        # precision limitation, reproduced on its own in
        # tests/test_payment_dispute_collision.py. It is not what these five
        # contracts are for, so the dispute window is left out of them.
        "Invoices shall be issued monthly in arrears unless the Statement of "
        "Work provides otherwise, and shall itemise the Services supplied "
        "during the billing period.",
    ])


def price_change(days: int) -> tuple[str, list[str]]:
    return ("Price Review", [
        f"The Supplier may revise the fees payable under this Agreement on not "
        f"less than {days} ({days}) days' prior written notice to the "
        "Customer. Any change in price takes effect from the start of the next "
        "billing period following expiry of that notice.",
        "A revision of the fees does not apply to Services already ordered "
        "under a Statement of Work that states a fixed price for a fixed term.",
    ])


def term_and_termination(convenience_days: int, cure_days: int,
                         purge_days: int, renewal_notice_days: int,
                         *, early_term_restriction: bool = True
                         ) -> tuple[str, list[str]]:
    paras = [
        "This Agreement commences on the Effective Date and continues for an "
        "initial committed term of twelve (12) months (the \"Initial Term\"), "
        "unless terminated earlier in accordance with this clause.",
        f"On expiry of the Initial Term this Agreement shall automatically "
        f"renew for successive periods of twelve (12) months, unless either "
        f"party gives written notice of non-renewal not less than "
        f"{renewal_notice_days} ({renewal_notice_days}) days before the end of "
        "the then-current term.",
        f"Either party may terminate this Agreement for convenience at any "
        f"time, for any reason, on not less than {convenience_days} "
        f"({convenience_days}) days' prior written notice to the other party.",
        f"Either party may terminate this Agreement immediately by written "
        f"notice if the other commits a material breach of this Agreement "
        f"which is not cured within {cure_days} ({cure_days}) days after "
        "receipt of written notice specifying the breach and requiring it to "
        "be remedied.",
        f"Within {purge_days} ({purge_days}) days after termination or expiry "
        "of this Agreement the Supplier shall securely purge all Customer data "
        "from its production systems, save for any copy it is required to "
        "retain by applicable law or that resides in routine backup media, "
        "which shall be purged in the ordinary course of the Supplier's backup "
        "cycle.",
    ]
    if early_term_restriction:
        paras.append(
            "The Customer has no right to terminate this Agreement before the "
            "expiry of the term stated in a Statement of Work that records a "
            "committed term, except for the Supplier's material breach. Where "
            "the Customer terminates such a Statement of Work early, an early "
            "termination fee equal to the charges for the remainder of the "
            "term becomes immediately due.")
    return ("Term and Termination", paras)


def confidentiality(survival_years: int) -> tuple[str, list[str]]:
    return ("Confidentiality", [
        "Each party shall keep the other party's Confidential Information "
        "strictly confidential, shall not disclose it to any third party "
        "except to those of its personnel and professional advisers who need "
        "to know it for the purposes of this Agreement, and shall not use it "
        "for any purpose other than the performance of this Agreement.",
        f"The obligations in this clause shall survive for a period of "
        f"{survival_years} ({survival_years}) years following the termination "
        "or expiry of this Agreement, after which each party's duty to protect "
        "Confidential Information under this Agreement ends, save in respect "
        "of any information that constitutes a trade secret, for which the "
        "obligations continue for so long as the information remains a trade "
        "secret under applicable law.",
        "The obligations in this clause do not apply to information that is or "
        "becomes public through no breach of this Agreement, that the "
        "receiving party already held free of any obligation of confidence, "
        "that is received from a third party entitled to disclose it, or that "
        "the receiving party independently develops without reference to the "
        "disclosing party's Confidential Information.",
    ])


def non_solicit(years: int) -> tuple[str, list[str]]:
    return ("Non-Solicitation", [
        f"During the term of this Agreement and for a period of {years} "
        f"({years}) years following its termination, neither party shall "
        "directly or indirectly solicit for employment any employee of the "
        "other party who was materially involved in the performance or receipt "
        "of the Services.",
        "This clause does not prevent either party from employing a person who "
        "responds to a general advertisement of employment not specifically "
        "targeted at the other party's personnel, or who approaches that party "
        "on their own initiative without any prior solicitation.",
    ])


def change_of_control(days: int) -> tuple[str, list[str]]:
    return ("Change of Control and Assignment", [
        f"Each party shall give the other not less than {days} ({days}) days' "
        "prior written notice of any change of control affecting it. For the "
        "purposes of this clause a change in control means a transaction or "
        "series of transactions as a result of which a person who did not "
        "previously control that party acquires control of it.",
        "Where a change of control results in the party coming under the "
        "control of a direct competitor of the other party, that other party "
        "may terminate this Agreement on thirty (30) days' written notice "
        "given within sixty (60) days of receiving notice of the change.",
    ])


def service_discontinuation_compound() -> tuple[str, list[str]]:
    """`AM-66` — the compound position: BOTH limbs, whichever is later."""
    return ("Service Discontinuation and End of Life", [
        "Where the Supplier decides to discontinue a Service, it shall give "
        "the Customer not less than thirty (30) days' advance written notice "
        "of the discontinuation, or shall continue to provide that Service "
        "until the end of the Customer's committed contract period, whichever "
        "is later.",
        "During the notice period the Supplier shall provide reasonable "
        "assistance to the Customer in migrating to an alternative service, "
        "including reasonable access to the Customer's data in a commonly "
        "used, machine-readable format.",
    ])


def service_discontinuation_notice_only() -> tuple[str, list[str]]:
    """The NARROWER form: the notice limb alone, without the committed term.

    `AM-66` is explicit that this is not the company position — the clause is
    narrower than the Constitution, so it must not confirm as a match.
    """
    return ("Service Discontinuation and End of Life", [
        "Where the Supplier decides to discontinue a Service, it shall give "
        "the Customer not less than thirty (30) days' advance written notice "
        "of the discontinuation.",
        "During the notice period the Supplier shall provide reasonable "
        "assistance to the Customer in migrating to an alternative service.",
    ])


def suspension_notice_cure() -> tuple[str, list[str]]:
    return ("Suspension of Services", [
        "The Supplier shall not suspend the Services except where the Customer "
        "has failed to pay an undisputed invoice when due, or where continued "
        "provision would breach applicable law or materially threaten the "
        "security or integrity of the Supplier's platform.",
        "Except where suspension is required immediately to protect the "
        "security of the platform or to comply with law, the Supplier shall "
        "give the Customer prior written notice before suspending the "
        "Services, together with a reasonable opportunity to cure the matter "
        "giving rise to the suspension.",
    ])


def governing_law_and_arbitration() -> tuple[str, list[str]]:
    return ("Governing Law and Dispute Resolution", [
        "This Agreement and any dispute or claim arising out of or in "
        "connection with it or its subject matter or formation, whether "
        "contractual or non-contractual, shall be governed by and construed in "
        "accordance with the laws of India, and the parties submit to the "
        "exclusive jurisdiction of the courts at Mumbai, Maharashtra.",
        "Any dispute arising out of or in connection with this Agreement that "
        "the parties cannot resolve through good-faith discussion within "
        "thirty (30) days shall be referred to and finally resolved by "
        "arbitration in accordance with the Arbitration and Conciliation Act, "
        "1996. The arbitration shall be conducted by a sole arbitrator "
        "appointed jointly by the parties, the seat of arbitration shall be "
        "Mumbai, and the language of the arbitration shall be English.",
    ])


def indemnity_and_ip() -> list[tuple[str, list[str]]]:
    return [
        ("Indemnification", [
            "The Supplier shall defend, indemnify, and hold harmless the "
            "Customer against any third-party claim alleging that the Services "
            "or the Deliverables, when used in accordance with this Agreement, "
            "infringe that third party's intellectual property rights, and "
            "shall pay any damages finally awarded or agreed in settlement.",
            "The Customer shall defend, indemnify, and hold harmless the "
            "Supplier against any third-party claim arising from the "
            "Customer's data or from the Customer's use of the Services in "
            "breach of this Agreement or of applicable law.",
            "The indemnified party shall notify the indemnifying party "
            "promptly of any claim, shall give it sole control of the defence "
            "and settlement, and shall provide reasonable cooperation at the "
            "indemnifying party's expense.",
        ]),
        ("Intellectual Property", [
            "Each party remains the sole and exclusive owner of all right, "
            "title, and interest in and to its own pre-existing intellectual "
            "property. Nothing in this Agreement transfers ownership of a "
            "party's background intellectual property to the other.",
            "All intellectual property rights in the Deliverables created "
            "specifically for the Customer under a Statement of Work vest in "
            "the Customer on payment in full, save that the Supplier retains "
            "ownership of any tools, libraries, methodologies and know-how of "
            "general application used in producing them, and grants the "
            "Customer a perpetual, non-exclusive licence to use those elements "
            "as embedded in the Deliverables.",
        ]),
    ]


# ---------------------------------------------------------------------------
# The five documents. `expected` records what each clause SHOULD produce, taken
# from the ratified standard's own number — the harness compares LegalMind's
# answer against this, so a disagreement is a finding either way round.
# ---------------------------------------------------------------------------
def _common_tail() -> list[tuple[str, list[str]]]:
    return [
        ("Representations and Warranties", FILLER["warranties"]),
        ("Compliance and Data Protection", FILLER["compliance"]),
        ("Force Majeure", FILLER["force_majeure"]),
        ("Notices and Relationship of the Parties", FILLER["notices"]),
        ("General", FILLER["general"]),
    ]


def contract_1() -> dict:
    """MSA-labelled, every position on the Constitution's number."""
    return {
        "slug": "C1-managed-services-all-match",
        "title": "MANAGED SERVICES AGREEMENT",
        "declared_type": "MSA",
        "parties": "NORTHWIND SYSTEMS PRIVATE LIMITED (the \"Supplier\") and "
                   "BRIGHTSHORE RETAIL PRIVATE LIMITED (the \"Customer\")",
        "sections": [
            ("Definitions", FILLER["definitions"]),
            ("Scope of Services", FILLER["scope"]),
            payment(21, 15),
            price_change(30),
            term_and_termination(30, 30, 30, 30),
            confidentiality(3),
            liability_capped(12),
            change_of_control(30),
            service_discontinuation_compound(),
            suspension_notice_cure(),
            *indemnity_and_ip(),
            governing_law_and_arbitration(),
            *_common_tail(),
        ],
        "expected": {
            "PAYMENT-PERIOD-MSA-001": "ACCEPTABLE",
            "PRICE-CHANGE-NOTICE-MSA-001": "ACCEPTABLE",
            "CONVENIENCE-NOTICE-MSA-001": "ACCEPTABLE",
            "CURE-PERIOD-MSA-001": "ACCEPTABLE",
            "DATA-PURGE-MSA-001": "ACCEPTABLE",
            "AUTORENEW-MSA-001": "ACCEPTABLE",
            "CONF-SURVIVAL-MSA-001": "ACCEPTABLE",
            "LIABILITY-MSA-001": "ACCEPTABLE",
            "CHANGE-OF-CONTROL-NOTICE-MSA-001": "ACCEPTABLE",
            "GOVLAW-MSA-001": "ACCEPTABLE",
            "ARBITRATION-MSA-001": "ACCEPTABLE",
            "INDEMNITY-MSA-001": "ACCEPTABLE",
            "IP-OWNERSHIP-MSA-001": "ACCEPTABLE",
            "GST-EXCLUSIVE-MSA-001": "ACCEPTABLE",
            "LIAB-EXCLUSIONS-MSA-001": "ACCEPTABLE",
            "EARLY-TERM-RESTRICTION-MSA-001": "ACCEPTABLE",
            "SERVICE-DISCONTINUATION-MSA-001": "ACCEPTABLE",
            "SUSPENSION-NOTICE-CURE-MSA-001": "ACCEPTABLE",
        },
    }


def contract_2() -> dict:
    """The headline test: the SAME positions under a non-MSA label.

    `AM-51`/`AM-60`: applicability follows what the document contains. If this
    scores differently from C1, the label is still deciding something.
    """
    c1 = contract_1()
    return {
        "slug": "C2-partner-agreement-all-match",
        "title": "STRATEGIC PARTNER AGREEMENT",
        "declared_type": "OTHER",
        "parties": "NORTHWIND SYSTEMS PRIVATE LIMITED (the \"Supplier\") and "
                   "CEDARLINE LOGISTICS PRIVATE LIMITED (the \"Partner\")",
        "sections": c1["sections"],
        "expected": dict(c1["expected"]),
    }


def contract_3() -> dict:
    """MSA-labelled, three positions deliberately off the number."""
    expected = dict(contract_1()["expected"])
    expected.update({
        "PAYMENT-PERIOD-MSA-001": "REQUIRES_MODIFICATION",     # 45 vs 21 days
        "CURE-PERIOD-MSA-001": "REQUIRES_MODIFICATION",        # 15 vs 30 days
        "CONF-SURVIVAL-MSA-001": "REQUIRES_MODIFICATION",      # 2 vs 3 years
    })
    return {
        "slug": "C3-technology-services-mixed",
        "title": "TECHNOLOGY SERVICES AGREEMENT",
        "declared_type": "MSA",
        "parties": "NORTHWIND SYSTEMS PRIVATE LIMITED (the \"Supplier\") and "
                   "HALEWOOD ANALYTICS PRIVATE LIMITED (the \"Customer\")",
        "sections": [
            ("Definitions", FILLER["definitions"]),
            ("Scope of Services", FILLER["scope"]),
            payment(45, 30),
            price_change(30),
            term_and_termination(30, 15, 30, 30),
            confidentiality(2),
            liability_capped(12),
            change_of_control(30),
            service_discontinuation_compound(),
            suspension_notice_cure(),
            *indemnity_and_ip(),
            governing_law_and_arbitration(),
            *_common_tail(),
        ],
        "expected": expected,
    }


def contract_4() -> dict:
    """OTHER-labelled, MSA and NDA positions in one document."""
    return {
        "slug": "C4-vendor-confidentiality-mixed",
        "title": "VENDOR SERVICES AND CONFIDENTIALITY AGREEMENT",
        "declared_type": "OTHER",
        "parties": "NORTHWIND SYSTEMS PRIVATE LIMITED (the \"Vendor\") and "
                   "ASHFIELD DIAGNOSTICS PRIVATE LIMITED (the \"Company\")",
        "sections": [
            ("Definitions", FILLER["definitions"]),
            ("Scope of Services", FILLER["scope"]),
            payment(21, 15),
            term_and_termination(30, 30, 30, 30),
            confidentiality(3),
            non_solicit(1),
            liability_capped(6),
            *indemnity_and_ip(),
            governing_law_and_arbitration(),
            *_common_tail(),
        ],
        "expected": {
            "PAYMENT-PERIOD-MSA-001": "ACCEPTABLE",
            "CONVENIENCE-NOTICE-MSA-001": "ACCEPTABLE",
            "CURE-PERIOD-MSA-001": "ACCEPTABLE",
            "CONF-SURVIVAL-MSA-001": "ACCEPTABLE",
            "GOVLAW-MSA-001": "ACCEPTABLE",
            "ARBITRATION-MSA-001": "ACCEPTABLE",
            "INDEMNITY-MSA-001": "ACCEPTABLE",
            "NON-SOLICIT-NDA-001": "REQUIRES_MODIFICATION",   # 1 vs 2 years
            "LIABILITY-MSA-001": "REQUIRES_MODIFICATION",     # 6 vs 12 months
        },
    }


def contract_5() -> dict:
    """An uncapped liability term — Constitution §9's Not Negotiable position."""
    return {
        "slug": "C5-distribution-needs-decision",
        "title": "DISTRIBUTION AND RESELLER AGREEMENT",
        "declared_type": "OTHER",
        "parties": "NORTHWIND SYSTEMS PRIVATE LIMITED (the \"Principal\") and "
                   "MERIDIAN CHANNEL PARTNERS PRIVATE LIMITED (the "
                   "\"Distributor\")",
        "sections": [
            ("Definitions", FILLER["definitions"]),
            ("Appointment and Scope", FILLER["scope"]),
            payment(21, 15),
            term_and_termination(30, 30, 30, 30),
            confidentiality(3),
            liability_uncapped(),
            service_discontinuation_notice_only(),
            *indemnity_and_ip(),
            governing_law_and_arbitration(),
            *_common_tail(),
        ],
        "expected": {
            "LIABILITY-MSA-001": "NEEDS_DECISION",   # uncapped -> §9 prohibited
            "PAYMENT-PERIOD-MSA-001": "ACCEPTABLE",
            "CURE-PERIOD-MSA-001": "ACCEPTABLE",
            "CONF-SURVIVAL-MSA-001": "ACCEPTABLE",
            "GOVLAW-MSA-001": "ACCEPTABLE",
            "ARBITRATION-MSA-001": "ACCEPTABLE",
            "INDEMNITY-MSA-001": "ACCEPTABLE",
        },
    }


CONTRACTS = [contract_1, contract_2, contract_3, contract_4, contract_5]


def write_docx(spec: dict, out_dir: Path) -> Path:
    import docx  # pip: python-docx

    document = docx.Document()
    document.add_heading(spec["title"], level=0)
    document.add_paragraph(
        "This Agreement is made between " + spec["parties"] + "."
    )
    document.add_paragraph(
        "SYNTHETIC TEST DOCUMENT — generated by tools/generate_test_contracts.py "
        "to exercise the analysis engine. Not a real agreement; no party named "
        "in it is real; it states no organisational legal position."
    )
    for index, (heading, paragraphs) in enumerate(spec["sections"], start=1):
        document.add_heading(f"{index}. {heading}", level=1)
        for number, text in enumerate(paragraphs, start=1):
            document.add_paragraph(f"{index}.{number}  {text}")

    document.add_heading("Execution", level=1)
    document.add_paragraph(
        "IN WITNESS WHEREOF the parties have executed this Agreement as of the "
        "Effective Date by their duly authorised representatives."
    )
    for role in ("Supplier", "Customer"):
        document.add_paragraph(f"Signed for and on behalf of the {role}:")
        document.add_paragraph("Name: ______________________")
        document.add_paragraph("Title: ______________________")
        document.add_paragraph("Date: ______________________")

    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{spec['slug']}.docx"
    document.save(path)
    return path


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--out", type=Path,
                    default=Path(__file__).resolve().parents[2]
                    / "legal-docs" / "synthetic",
                    help="output directory (gitignored; never the repository)")
    args = ap.parse_args(argv)

    import json
    manifest = {}
    for builder in CONTRACTS:
        spec = builder()
        path = write_docx(spec, args.out)
        words = sum(len(p.split())
                    for _, paras in spec["sections"] for p in paras)
        manifest[spec["slug"]] = {
            "title": spec["title"],
            "declared_type": spec["declared_type"],
            "file": path.name,
            "expected": spec["expected"],
        }
        print(f"{path.name:38} {spec['declared_type']:5} "
              f"{words:5} words  ~{max(1, round(words / 500))} pages  "
              f"{len(spec['expected'])} expected outcomes")

    manifest_path = args.out / "expected.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"\nmanifest: {manifest_path}")
    return 0


if __name__ == "__main__":                            # pragma: no cover
    raise SystemExit(main())

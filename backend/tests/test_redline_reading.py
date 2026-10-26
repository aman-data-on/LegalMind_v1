"""Marked-up (redlined) PDFs — AM-146.

A real customer MOU, exported with its tracked changes shown, was read as plain text:
a struck "3" beside an inserted "5" became "53 years", the markup tool's "[13.]" label
glued clause 13 onto clause 12, and bold headings above short bullets were absorbed
into the section above. Each test rebuilds one of those on a synthetic PDF (54.6 keeps
real documents out of the repository): orange text, an underline drawn under an
insertion and a line drawn through a deletion, as Word and LibreOffice export them.
"""

from __future__ import annotations

import pytest

from legalmind.ingestion import parsing
from legalmind.ingestion.parsing import detect_clause_number, parse_pdf

pymupdf = pytest.importorskip("pymupdf")

ORANGE = (0.776, 0.573, 0.0)
BLACK = (0, 0, 0)
BLUE = (0.0, 0.3, 0.8)
GREEN = (0.0, 0.5, 0.2)


class _Page:
    """Writes runs left to right on one line; marks a run struck or underlined."""

    def __init__(self, page, bars: bool = False):
        self.page, self.bars = page, bars

    def line(self, y: float, *runs: tuple[str, str], x: float = 72.0) -> None:
        for text, mark in runs:
            font = "hebo" if mark == "bold" else "helv"
            colour = (ORANGE if mark in ("ins", "del") else GREEN if mark == "ins2" else
                      BLUE if mark in ("link", "anchor") else BLACK)
            self.page.insert_text((x, y), text, fontsize=10, fontname=font, color=colour)
            end = x + pymupdf.get_text_length(text, fontname=font, fontsize=10)
            # Each rule stays inside its own glyphs, as the exporters draw it (the real
            # MOU: an underline ends at 293.80, the next word starts at 293.85).
            if mark in ("ins", "ins2", "link", "anchor", "black_underline"):
                self.page.draw_line((x + .05, y + 2), (end - .05, y + 2), color=colour,
                                    width=0.5)
            elif mark == "del":
                self.page.draw_line((x + .05, y - 3.5), (end - .05, y - 3.5),
                                    color=colour, width=0.5)
            if mark == "anchor":
                self.page.insert_link({"kind": pymupdf.LINK_URI, "uri": "https://x.test",
                                       "from": pymupdf.Rect(x, y - 9, end, y + 3)})
            if self.bars and mark in ("ins", "ins2", "del"):     # the changed-line mark
                self.page.draw_line((28, y - 9), (28, y + 3), color=BLACK, width=0.5)
            x = end


def _pdf(*pages: list[tuple], bars: bool = False) -> bytes:
    doc = pymupdf.open()
    for lines in pages:
        page = _Page(doc.new_page(), bars)
        for y, runs in lines:
            page.line(y, *runs)
    return doc.tobytes()


def _read(data: bytes):
    return parse_pdf(data, ocr=False, docling=False).segments


def test_a_struck_figure_beside_an_inserted_one_is_two_versions_never_one_number():
    segments = _read(_pdf([
        (72, [("Confidentiality", "bold")]),
        (90, [("Confidentiality obligation continues for ", ""), ("5", "ins"),
              ("3", "del"), (" years after termination.", "")]),
    ]))
    clause = segments[-1]
    assert "continues for 5 years" in clause.content
    assert "53" not in clause.content
    revision = clause.metadata["revision"]
    assert "continues for 3 years" in revision["original"]
    assert revision["changes"] == [{"deleted": "3", "inserted": "5"}]


def test_a_deleted_sentence_and_its_replacement_read_as_the_proposed_text():
    segments = _read(_pdf([
        (72, [("13. Service Termination: ", ""),
              ("A 30-day prior notice period applies.", "del"), (" ", ""),
              ("The Customer may terminate on 15 days notice.", "ins")]),
    ]))
    clause = segments[0]
    assert "15 days notice" in clause.content and "30-day" not in clause.content
    assert "A 30-day prior notice period applies." in clause.metadata["revision"]["original"]
    assert "15 days" not in clause.metadata["revision"]["original"]


def test_a_wholly_deleted_bullet_leaves_the_proposed_text_and_stays_in_the_original():
    segments = _read(_pdf([
        (72, [("Termination", "bold")]),
        (90, [("- Services stop after termination", "")]),
        (104, [("- ", ""), ("Early termination charges may apply", "del")]),
        (118, [("- ", ""), ("Customer must give 30 days notice", "del")]),
        (132, [("- Outstanding dues must be cleared", "")]),
    ]))
    clause = segments[-1]
    assert "Early termination" not in clause.content
    assert clause.content.count("- ") == 2              # no bullet left standing empty
    original = clause.metadata["revision"]["original"]
    assert original.index("Early termination") < original.index("30 days notice")


def test_a_document_with_links_but_no_strikes_is_not_marked_up():
    data = _pdf([
        (72, [("1. Contact ", ""), ("corporate@example.com", "link"),
              (" for notices.", "")]),
        (90, [("Ref: ", ""), ("LS/OCT/2026/NK/0156", "black_underline")]),
    ])
    doc = pymupdf.open(stream=data, filetype="pdf")
    assert parsing._revision_colours(doc) == frozenset()
    assert not any("revision" in s.metadata for s in _read(data))


def test_an_unmarked_document_reads_exactly_as_without_the_redline_reader(monkeypatch):
    # The isolation the owner asked for: the gate is the ONLY way in. Bold headings
    # above short lines keep today's reading on a document with no tracked changes.
    data = _pdf([
        (72, [("Confidentiality", "bold")]),
        (90, [("- Both parties keep information confidential", "")]),
        (104, [("Governing Law", "bold")]),
        (118, [("- Governed by Indian law", "")]),
    ])
    with_gate = [(s.content, s.metadata) for s in _read(data)]
    monkeypatch.setattr(parsing, "_revision_colours", lambda doc: frozenset())
    assert with_gate == [(s.content, s.metadata) for s in _read(data)]
    # AM-196 (2026-10-26): a heading over a sentence bullet now starts a section on EVERY
    # document — the gate's isolation above is unchanged, the plain reading improved
    assert len(with_gate) == 2


def test_in_a_marked_up_document_a_bold_heading_above_short_bullets_starts_a_section():
    segments = _read(_pdf([
        (72, [("Confidentiality", "bold")]),
        (90, [("- Obligation continues for ", ""), ("5", "ins"), ("3", "del"),
              (" years", "")]),
        (104, [("Governing Law", "bold")]),
        (118, [("- Governed by Indian law", "")]),
        (132, [("Penalty Clause:", "bold")]),
        (146, [("Credits are paid monthly.", "")]),
    ]))
    titles = [s.section_title for s in segments]
    assert titles == ["Confidentiality", "Governing Law", "Penalty Clause"]
    assert all(s.metadata.get("heading") for s in segments)


def test_a_markup_tools_bracketed_number_is_the_clause_number():
    assert detect_clause_number("[13.] Service Termination: notice")[0] == "13"
    assert detect_clause_number("[4.2.] Payment")[0] == "4.2"
    assert detect_clause_number("[1] See the schedule")[0] is None      # a footnote
    assert detect_clause_number("[13.] the rest of a sentence")[0] is None


def test_a_bracketed_clause_number_no_longer_glues_onto_the_clause_above():
    segments = _read(_pdf([
        (72, [("12. Support is provided on tickets raised by the customer.", "")]),
        (90, [("[13.] Service Termination: ", ""),
              ("The Customer may terminate on 15 days notice.", "ins")]),
    ]))
    assert [s.section_number for s in segments] == ["12", "13"]
    assert "15 days" not in segments[0].content


def test_a_clause_running_onto_the_next_page_is_one_segment():
    segments = _read(_pdf(
        [(72, [("Penalty Clause:", "bold")]),
         (90, [("Credits are calculated on a pro-rata basis from the time", ""),
               (" of the ticket", "ins"), (" daily", "del")])],
        [(72, [("the Trouble Ticket is issued until resolution.", "")]),
         (90, [("Dispute Resolution:", "bold")]),
         (104, [("Courts at Pune have jurisdiction.", "")])],
    ))
    penalty = next(s for s in segments if s.section_title == "Penalty Clause")
    assert "from the time of the ticket\nthe Trouble Ticket is issued" in penalty.content
    assert (penalty.page_number, penalty.metadata["end_page"]) == (1, 2)
    assert any(s.section_title == "Dispute Resolution" for s in segments)


def test_a_capitalised_page_top_after_an_unpunctuated_line_is_not_a_continuation():
    # A signature block ends unpunctuated; "IN WITNESS WHEREOF" begins its own text.
    segments = _read(_pdf(
        [(72, [("Warm Regards", "")]), (90, [("Business Development Manager", ""),
                                            (" Ltd", "del")])],
        [(72, [("IN WITNESS WHEREOF the Parties have executed this Agreement.", "")])],
    ))
    assert segments[-1].page_number == 2
    assert "end_page" not in segments[0].metadata


def test_a_docx_tracked_change_keeps_its_original_beside_the_accepted_text():
    """The DOCX path read the accepted view already (A3) but threw the deleted text
    away; it is now recorded exactly as a PDF's is."""
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    w = nsdecls("w")
    d = docx.Document()
    p = d.add_paragraph("Liability shall not exceed the total fees paid in the ")
    for xml in (f'<w:del {w} w:id="1" w:author="x"><w:r><w:delText>12</w:delText></w:r></w:del>',
                f'<w:ins {w} w:id="2" w:author="x"><w:r><w:t>6</w:t></w:r></w:ins>',
                f'<w:r {w}><w:t xml:space="preserve"> months.</w:t></w:r>'):
        p._p.append(parse_xml(xml))
    d.add_paragraph("Governing law is India.")
    import io
    buf = io.BytesIO()
    d.save(buf)
    segments = parsing.parse_docx(buf.getvalue()).segments
    clause, plain = segments[0], segments[1]
    assert "paid in the 6 months" in clause.content
    assert clause.metadata["revision"] == {
        "original": "Liability shall not exceed the total fees paid in the 12 months.",
        "changes": [{"deleted": "12", "inserted": "6"}]}
    assert "revision" not in plain.metadata


# --- AM-157 (2026-10-17): nothing the client marked is missed --------------------------

def _changes(segments) -> list[dict]:
    return [c for s in segments for c in s.metadata.get("revision", {}).get("changes", ())]


def test_a_client_who_only_added_underlined_text_is_read_as_marked_up():
    # No strike anywhere: before AM-157 the whole document read as unmarked and
    # none of its additions was recorded.
    segments = _read(_pdf([
        (72, [("1. Fees are payable monthly", ""), (" subject to prior approval", "ins")]),
        (90, [("2. Payment is due within ", ""), ("1", "ins"), ("5 days.", "")]),
    ]))
    assert _changes(segments) == [{"deleted": "", "inserted": "subject to prior approval"},
                                  {"deleted": "", "inserted": "1"}]
    assert "within 15 days" in segments[-1].content


def test_a_wholly_new_clause_beside_a_change_bar_is_an_insertion():
    segments = _read(_pdf([
        (72, [("Notices", "bold")]), (90, [("- Notices are given in writing", "")]),
        (110, [("Non-Solicitation", "ins")]),
        (126, [("- Neither party shall solicit staff", "ins")]),
    ], bars=True))
    assert " ".join(c["inserted"] for c in _changes(segments)) == \
        "Non-Solicitation - Neither party shall solicit staff"


def test_coloured_underlined_lines_alone_are_a_template_style_not_markup():
    # The limit of AM-157, and the isolation it keeps: whole lines in a coloured
    # underline, with no strike, no plain text beside them and no change bar, are how
    # a template styles a heading — the document reads exactly as before.
    data = _pdf([(72, [("Non-Solicitation", "ins")]),
                 (90, [("- Neither party shall solicit staff", "")])])
    assert parsing._revision_colours(pymupdf.open(stream=data, filetype="pdf")) == frozenset()


def test_an_underlined_link_inline_with_text_is_not_an_insertion():
    data = _pdf([(72, [("1. See our ", ""), ("Terms of Service", "anchor"),
                       (" and write to ", ""), ("legal@example.com", "link"), (".", "")])])
    assert parsing._revision_colours(pymupdf.open(stream=data, filetype="pdf")) == frozenset()


def test_a_second_reviewer_who_only_inserts_is_read_too():
    segments = _read(_pdf([
        (72, [("1. Notice period is ", ""), ("30", "del"), ("15", "ins"), (" days.", "")]),
        (90, [("2. Courts at Pune", ""), (" and Bengaluru", "ins2"), (" decide.", "")]),
    ]))
    assert _changes(segments) == [{"deleted": "30", "inserted": "15"},
                                  {"deleted": "", "inserted": "and Bengaluru"}]


def test_a_deletion_then_an_insertion_never_trade_letters():
    segments = _read(_pdf([(72, [
        ("1. Courts at ", ""), ("Mumbai", "del"), ("Bangalore", "ins"), (" and seat at ", ""),
        ("Mumbai", "del"), ("Bengaluru", "ins"), (" for ", ""), ("3", "del"), ("5", "ins"),
        (" years.", "")])]))
    assert "Bangalore and seat at Bengaluru for 5 years" in segments[0].content
    assert _changes(segments) == [{"deleted": "Mumbai", "inserted": "Bangalore"},
                                  {"deleted": "Mumbai", "inserted": "Bengaluru"},
                                  {"deleted": "3", "inserted": "5"}]


@pytest.mark.parametrize("pages", [
    [[(72, [("1. Liability is capped ", ""), ("except for fraud, wilful", "ins")]),
      (86, [("misconduct or gross negligence", "ins"), (", at fees paid.", "")])]],
    [[(72, [("1. Liability is capped ", ""), ("except for fraud, wilful", "ins")])],
     [(72, [("misconduct or gross negligence", "ins"), (", at fees paid.", "")])]],
], ids=["over two lines", "over a page break"])
def test_an_insertion_that_runs_on_keeps_every_word(pages):
    words = " ".join(c["inserted"] for c in _changes(_read(_pdf(*pages))))
    assert words == "except for fraud, wilful misconduct or gross negligence"


def test_a_zero_width_ligature_glyph_takes_the_mark_of_the_change_it_sits_in():
    # The real MOU's struck "confirms": its "fi" ligature left a zero-width "i" that no
    # strike crosses, and it was read as inserted ("…amended terms.i").
    def span(text, flags, x0, x1):
        return {"text": text, "color": 1, "char_flags": flags, "bbox": (x0, 0, x1, 9)}
    styled = {"blocks": [{"bbox": (0, 0, 400, 9), "number": 0, "lines": [{"spans": [
        span("Acceptance is by signature.", 2, 0, 150), span("Use confirms conf", 1, 150, 240),
        span("i", 0, 240, 240), span("rms it", 1, 240, 270)]}]}]}
    marked = parsing._revision_blocks(styled, frozenset({1}))[0][4]
    proposed, changes = parsing._resolve_revisions(marked)
    assert proposed.strip() == "Acceptance is by signature."
    assert [(c.deleted, c.inserted) for c in changes] == [
        ("Use confirms confirms it", "Acceptance is by signature.")]


def _docx(*paragraphs: list[str], table: list[str] | None = None) -> list:
    """Paragraphs of raw run XML; `table`: one row's cell XML, appended after them."""
    import io

    import docx
    from docx.oxml import parse_xml
    d = docx.Document()
    for runs in paragraphs:
        p = d.add_paragraph()
        for xml in runs:
            p._p.append(parse_xml(xml))
    if table:
        cells = d.add_table(rows=1, cols=len(table)).rows[0].cells
        for cell, xml in zip(cells, table, strict=True):
            for run in parse_xml(f"<w:p {_W}>{xml}</w:p>"):
                cell.paragraphs[0]._p.append(run)
    buf = io.BytesIO()
    d.save(buf)
    return parsing.parse_docx(buf.getvalue()).segments


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _run(text: str) -> str:
    return f'<w:r {_W}><w:t xml:space="preserve">{text}</w:t></w:r>'


def _mark(tag: str, text: str, n: int) -> str:
    element = "w:delText" if tag in ("del", "moveFrom") else "w:t"
    return (f'<w:{tag} {_W} w:id="{n}" w:author="c" w:name="m"><w:r><{element} '
            f'xml:space="preserve">{text}</{element}></w:r></w:{tag}>')


def test_a_docx_paragraph_deleted_whole_is_kept_on_its_neighbour_never_lost():
    # The Word copy of the real MOU lost four deleted paragraphs this way (40% of its
    # deleted text) — they had no accepted text to hang on.
    segments = _docx([_mark("del", "Opening recital removed.", 1)], [_run("Clause one stays.")],
                     [_mark("del", "Early termination charges apply.", 2)])
    assert [s.content for s in segments] == ["Clause one stays."]
    assert segments[0].metadata["revision"] == {
        "original": "Opening recital removed.\nClause one stays.\n"
                    "Early termination charges apply.",
        "changes": [{"deleted": "Opening recital removed.", "inserted": ""},
                    {"deleted": "Early termination charges apply.", "inserted": ""}]}


def test_a_docx_tracked_change_in_a_table_cell_is_recorded():
    segments = _docx([_run("Commercials")], table=[
        _run("Notice period"),
        _run("Notice: ") + _mark("del", "30", 3) + _mark("ins", "15", 4) + _run(" days")])
    table = segments[-1]
    assert "Notice: 15 days" in table.content
    assert table.metadata["revision"]["changes"] == [{"deleted": "30", "inserted": "15"}]
    assert "Notice: 30 days" in table.metadata["revision"]["original"]


def test_a_docx_moved_paragraph_is_a_deletion_where_it_was_and_an_insertion_where_it_is():
    segments = _docx([_mark("moveFrom", "Governing law is India.", 5)],
                     [_run("Notices are in writing.")],
                     [_mark("moveTo", "Governing law is India.", 6)])
    assert [s.content for s in segments] == ["Notices are in writing.",
                                             "Governing law is India."]
    assert _changes(segments) == [{"deleted": "Governing law is India.", "inserted": ""},
                                  {"deleted": "", "inserted": "Governing law is India."}]


# --- AM-157 (2026-10-17): tables, images and scans ----------------------------------

needs_tesseract = pytest.mark.skipif(not parsing.ocr_available(),
                                     reason="needs the OCR toolchain")


def _table_pdf(rows: list[list[list[tuple[str, str]]]]) -> bytes:
    """A ruled table, one cell per (runs) list, under a marked opening line."""
    doc = pymupdf.open()
    page = _Page(doc.new_page())
    page.line(72, ("1. The schedule applies ", ""), ("in full", "ins"))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            box = pymupdf.Rect(72 + j * 200, 100 + i * 22, 272 + j * 200, 122 + i * 22)
            page.page.draw_rect(box, color=BLACK, width=0.7)
            page.line(box.y0 + 15, *cell, x=box.x0 + 4)
    return doc.tobytes()


def test_a_tracked_change_in_a_table_cell_keeps_the_table_and_both_versions():
    segments = _read(_table_pdf([
        [[("Service", "")], [("Notice period", "")]],
        [[("Termination", "")], [("30", "del"), ("15", "ins"), (" days", "")]],
        [[("Tier 2", "del")], [("10%", "del")]],
        [[("Tier 3", "ins")], [("20%", "ins")]]]))
    text = "\n".join(s.content for s in segments)
    assert "| Termination | 15 days |" in text and "3015" not in text
    assert "Tier 2" not in text and "| Tier 3 | 20% |" in text
    assert {"deleted": "30", "inserted": "15"} in _changes(segments)


def test_a_strike_is_told_from_a_letters_own_bar():
    import numpy as np
    # A strike: free at both ends, above and below, letters inside it.
    band = np.zeros((1, 20), bool)
    band[0, 2:14] = True
    above = np.zeros(20, bool)
    above[5:11] = True
    below = above.copy()
    assert parsing._strikes(band, above, below, 8, 30)[2:14].all()
    # An "A"'s bar: free above, but its legs spread below and meet the ends.
    below[2:14] = True
    assert not parsing._strikes(band, above, below, 8, 30).any()
    # A long run through several letters is a strike whatever its ends meet.
    assert parsing._strikes(band, above, below, 8, 10)[2:14].all()


def test_a_change_running_across_segments_is_split_never_dropped():
    # OCR's paragraph break inside an insertion: the change was dropped whole.
    text = "Liability is capped except for fraud\n\nor gross negligence at fees."
    segments = parsing.segment_paragraphs(text, page_number=1, source_type=parsing.
                                          EvidenceSourceType.OCR, base_offset=0)
    start = text.index("except")
    end = text.index(" at fees")
    placed = parsing._record_revisions(segments, [parsing._Change(
        start, end, "", text[start:end])], 0)
    assert placed == 2
    assert " ".join(c["inserted"] for c in _changes(segments)) == \
        "except for fraud or gross negligence"


def test_a_docx_picture_that_cannot_be_read_makes_the_reading_partial(monkeypatch):
    import io

    import docx
    from docx.shared import Inches
    monkeypatch.setattr(parsing, "ocr_available", lambda: False)
    monkeypatch.setattr(parsing, "_tesseract_pixmap",
                        lambda pix: (_ for _ in ()).throw(FileNotFoundError("tesseract")))
    d = docx.Document()
    d.add_paragraph("1. Fees are payable monthly.")
    picture = pymupdf.open().new_page(width=400, height=200).get_pixmap()
    d.add_picture(io.BytesIO(picture.tobytes("png")), width=Inches(6))
    d.add_picture(io.BytesIO(picture.tobytes("png")), width=Inches(1))    # a logo
    buf = io.BytesIO()
    d.save(buf)
    result = parsing.parse_docx(buf.getvalue())
    assert result.status == parsing.ExtractionStatus.PARTIAL
    assert result.diagnostics.count("a picture not read (OCR toolchain unavailable)") == 1


def _scan(*lines: list[tuple], colour=None) -> bytes:
    """An image-only PDF of the given `_Page` lines, as a 200-dpi colour scan."""
    src = pymupdf.open()
    page = _Page(src.new_page())
    for y, runs in lines:
        page.line(y, *runs)
    if colour:
        page.page.insert_text((72, 300), "HOME   CONTACT   New   Login", fontsize=14,
                              fontname="hebo", color=colour)
    scan = pymupdf.open()
    out = scan.new_page(width=page.page.rect.width, height=page.page.rect.height)
    out.insert_image(out.rect, stream=page.page.get_pixmap(dpi=200).tobytes("png"))
    return scan.tobytes()


@needs_tesseract
def test_a_scanned_redline_is_read_as_two_versions_never_merged():
    result = parse_pdf(_scan(
        (100, [("Confidentiality continues for ", ""), ("5", "ins"), ("3", "del"),
               (" years after termination.", "")]),
        (130, [("Disputes are resolved by arbitration in ", ""), ("Bangalore", "ins"),
               ("Mumbai", "del"), (", India.", "")])), ocr=True, docling=False)
    text = " ".join(s.content for s in result.segments)
    assert "for 5 years" in text and "53" not in text
    assert "in Bangalore, India" in text and "Mumbai" not in text
    assert {"deleted": "3", "inserted": "5"} in _changes(result.segments)


@needs_tesseract
def test_a_clean_scan_with_coloured_menus_is_not_marked_up():
    # A web print's coloured "HOME" / "CONTACT" (their H and A bars) and a bold "New"
    # read as struck before the strike geometry (2026-10-17).
    result = parse_pdf(_scan(
        (100, [("1. Fees are payable monthly in advance.", "")]),
        (130, [("2. Governing law is India.", "")]), colour=(0.85, 0.15, 0.2)),
        ocr=True, docling=False)
    assert not _changes(result.segments)


@needs_tesseract
def test_a_docx_picture_in_the_text_is_read_by_ocr_and_a_logo_is_not():
    import io

    import docx
    from docx.shared import Inches
    page = pymupdf.open().new_page(width=450, height=250)
    for i, line in enumerate(["Schedule A - Service Credits", "Tier 1: 5% of monthly fee",
                              "Liability capped at 6 months of fees"]):
        page.insert_text((15, 40 + i * 50), line, fontsize=18, fontname="helv")
    png = page.get_pixmap(dpi=200).tobytes("png")
    d = docx.Document()
    d.add_paragraph("1. Fees are payable monthly.")
    d.add_picture(io.BytesIO(png), width=Inches(6.2))
    d.add_picture(io.BytesIO(png), width=Inches(1.2))                     # logo-sized
    buf = io.BytesIO()
    d.save(buf)
    result = parsing.parse_docx(buf.getvalue())
    read = [s for s in result.segments if s.source_type == parsing.EvidenceSourceType.OCR]
    assert len(read) == 1 and "6 months" in read[0].content
    assert result.status == parsing.ExtractionStatus.COMPLETE


def test_docling_tables_are_not_used_on_a_marked_up_page():
    # Docling reads a table's text layer, where a struck "30" and an inserted "15" are
    # "3015": on a marked-up page the table is built from the marked text instead.
    data = _pdf([(72, [("1. Notice is ", ""), ("30", "del"), ("15", "ins"), (" days.", "")]),
                 (90, [("Credits | 5%", "")])])
    doc = pymupdf.open(stream=data, filetype="pdf")
    page, colours = doc[0], parsing._revision_colours(doc)
    revision = (parsing._styled_text(page), colours)
    table = {"bbox": tuple(page.rect), "rows": [["1. Notice is 3015 days.", "Credits | 5%"]],
             "cells": [{"text": "1. Notice is 3015 days.", "bbox": tuple(page.rect)}]}
    text, by_docling = parsing._page_blocks_in_reading_order(page, [table], revision)
    assert not by_docling
    assert "Notice is 15 days" in parsing._resolve_revisions(text)[0]


# --- Review of PR #214 (2026-10-17) ---------------------------------------------------

def test_a_merged_change_across_a_paragraph_break_keeps_every_letter():
    # `_resolve_revisions` joins runs across "\n\n" with ONE space, so the change's
    # offsets and its text no longer line up; its parts are read from the text itself.
    marked = (f"Liability is capped {parsing._INS_OPEN}except for fraud{parsing._INS_CLOSE}"
              f"\n\n{parsing._INS_OPEN}or gross negligence{parsing._INS_CLOSE} at fees.")
    proposed, changes = parsing._resolve_revisions(marked)
    segments = parsing.segment_paragraphs(proposed, page_number=1, source_type=parsing.
                                          EvidenceSourceType.OCR, base_offset=0)
    parsing._record_revisions(segments, changes, 0)
    assert [c["inserted"] for c in _changes(segments)] == ["except for fraud",
                                                          "or gross negligence"]


def test_a_blue_underline_in_a_sentence_is_a_link_without_its_annotation():
    # "Microsoft Print to PDF" drops link annotations: a blue underlined policy name in
    # a clean contract made the whole document "marked up".
    data = _pdf([(72, [("1. Use is governed by our ", ""), ("Acceptable Use Policy", "link"),
                       (" as published.", "")])])
    assert parsing._revision_colours(pymupdf.open(stream=data, filetype="pdf")) == frozenset()


def test_a_docx_page_size_without_its_attributes_does_not_fail_the_reading():
    import io

    import docx
    d = docx.Document()
    d.add_paragraph("1. Fees are payable monthly.")
    for attribute in ("w", "h"):
        d.sections[0]._sectPr.pgSz.attrib.pop(parsing._W + attribute, None)
    buf = io.BytesIO()
    d.save(buf)
    assert parsing.parse_docx(buf.getvalue()).segments[0].content == \
        "1. Fees are payable monthly."

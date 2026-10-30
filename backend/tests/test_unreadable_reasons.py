"""QA-032/034/037 (2026-10-30): a file that cannot be read fails as a ParseError with a
reason the screens can say — never a 500, never a Python exception name."""

from __future__ import annotations

import io
import zipfile

import pytest

from legalmind.ingestion import parsing
from legalmind.ingestion.parsing import DOCX_MIME, PDF_MIME


def _docx(mutate) -> bytes:
    import docx
    document = docx.Document()
    document.add_paragraph("1. Services")
    document.add_table(rows=1, cols=2).cell(0, 0).text = "x"
    raw = io.BytesIO()
    document.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as src, \
            zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(item, mutate(item.filename, src.read(item.filename)))
    return out.getvalue()


def _pdf(pages: int, password: str | None = None) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for _ in range(pages):
        doc.new_page().insert_text((72, 72), "1. Services")
    if password:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw=password,
                           owner_pw=password)
    return doc.tobytes()


def _code(data: bytes, mime: str) -> str:
    with pytest.raises(parsing.ParseError) as caught:
        parsing.parse(data, mime)
    return caught.value.code


def test_a_malformed_table_attribute_is_a_damaged_file_not_a_crash():
    data = _docx(lambda name, body: body.replace(
        b"<w:tcPr>", b'<w:tcPr><w:gridSpan w:val="abc"/>')
        if name == "word/document.xml" else body)
    assert _code(data, DOCX_MIME) == "CORRUPT"


def test_a_macro_enabled_word_file_says_so():
    data = _docx(lambda name, body: body.replace(
        b"wordprocessingml.document.main+xml", b"ms-word.document.macroEnabled.main+xml")
        if name == "[Content_Types].xml" else body)
    assert _code(data, DOCX_MIME) == "MACRO_ENABLED"


def test_a_password_protected_pdf_says_so():
    assert _code(_pdf(1, password="u"), PDF_MIME) == "PASSWORD_PROTECTED"


def test_a_damaged_pdf_says_so():
    assert _code(b"%PDF-1.7 not really", PDF_MIME) == "CORRUPT"


def test_a_pdf_over_the_page_cap_is_refused_before_any_page_is_read(monkeypatch):
    monkeypatch.setattr(parsing, "MAX_PDF_PAGES", 3)
    assert _code(_pdf(4), PDF_MIME) == "TOO_MANY_PAGES"
    assert parsing.parse(_pdf(3), PDF_MIME).pages_total == 3


def test_a_docx_with_a_doctype_is_refused_not_read_with_text_missing():
    """QA-096: an entity declared in a DOCTYPE swallowed the text after it, silently."""
    data = _docx(lambda name, body: body.replace(
        b"<w:document", b'<!DOCTYPE w:document [<!ENTITY x "y">]><w:document', 1)
        if name == "word/document.xml" else body)
    assert _code(data, DOCX_MIME) == "CORRUPT"

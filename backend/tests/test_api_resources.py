"""Contracts, documents, reviews, configuration, report and administration.

Locked references: 49.3, 49.6, 49.8, Step 29, Step 34, 42.4, 42.12, ENG-09,
F-8, F-9, S-8, S-9, S-10, SEC-05.
"""

from __future__ import annotations

import io
import json
import uuid

import pytest
from sqlalchemy import func, select

from legalmind.db import models as M
from legalmind.domain import enums as E
from legalmind.security import permissions as P
from tests.conftest import (
    bespoke_role,
    grant,
    grant_role,
    make_review_for,
    make_user,
    sign_in,
    sign_out,
)

V1 = "/api/v1"

DOCX_MIME = ("application/vnd.openxmlformats-officedocument"
             ".wordprocessingml.document")


@pytest.fixture
def owner(db, seeded):
    user = make_user(db)
    grant_role(db, user, P.ROLE_USER)
    return user


def build_docx(paragraphs: list[str]) -> bytes:
    import docx
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# =====================================================================
# Contracts
# =====================================================================
def test_create_read_and_patch_a_contract(api, db, owner):
    sign_in(api, db, owner)
    created = api.post(f"{V1}/contracts",
                       json={"name": "ACME MSA", "contract_type": "MSA"})
    assert created.status_code == 201
    contract_id = created.json()["data"]["id"]
    assert created.json()["data"]["owner_id"] == str(owner.id)
    assert created.json()["data"]["status"] == "DRAFT"

    patched = api.patch(f"{V1}/contracts/{contract_id}",
                        json={"status": "ACTIVE"})
    assert patched.json()["data"]["status"] == "ACTIVE"
    assert api.get(f"{V1}/contracts/{contract_id}"
                   ).json()["data"]["name"] == "ACME MSA"


def test_put_is_not_accepted_on_a_contract(api, db, owner):
    """49.1 — PUT is not used, so it is not routed."""
    sign_in(api, db, owner)
    created = api.post(f"{V1}/contracts", json={"name": "ACME MSA"})
    contract_id = created.json()["data"]["id"]
    assert api.request("PUT", f"{V1}/contracts/{contract_id}",
                       json={"name": "x"}).status_code in {404, 405}


# =====================================================================
# Document upload and download — Step 34
# =====================================================================
def test_upload_produces_a_document_version_and_evidence(api, db, owner):
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]

    response = api.post(
        f"{V1}/contracts/{contract_id}/document-versions",
        content=build_docx(["1. Limitation of Liability",
                            "Liability is capped at fees paid."]),
        headers={"Content-Type": DOCX_MIME, "X-Filename": "msa.docx"})
    assert response.status_code == 201
    body = response.json()["data"]
    assert body["document_version"]["version_number"] == 1
    assert body["document_version"]["processing_status"] == "COMPLETED"
    assert body["evidence_count"] >= 1
    assert body["duplicate_of"] is None
    # Storage coordinates are never part of the resource.
    assert "storage_key" not in body["document_version"]


def test_a_non_ascii_filename_round_trips_through_the_percent_encoded_header(
        api, db, owner):
    """2026-09-03 defect, reported against the live site with a real filename.

    HTTP header values are restricted to ISO-8859-1 bytes. A filename with any
    character outside that range — an en dash, a curly quote, anything non-
    ASCII — made the browser's `fetch` throw before the request even left the
    tab, so the upload never reached the server at all: no request logged, no
    structured error, just the client's generic fallback message. The frontend
    now percent-encodes `X-Filename`; this pins the server's half — decoding it
    back to the real filename rather than storing the encoded form."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]

    # The en dash (U+2013) is the POINT of this test — a real uploaded filename
    # contained one and threw in `fetch` before reaching the network. RUF001
    # flags ambiguous characters in strings; here the ambiguity is the subject,
    # so the character must not be "corrected".
    real_name = "MUTUAL NON – DISCLOSURE AGREEMENT.docx"  # noqa: RUF001 - en dash, U+2013
    from urllib.parse import quote
    response = api.post(
        f"{V1}/contracts/{contract_id}/document-versions",
        content=build_docx(["1. Limitation of Liability",
                            "Liability is capped at fees paid."]),
        headers={"Content-Type": DOCX_MIME, "X-Filename": quote(real_name)})
    assert response.status_code == 201, response.text
    assert response.json()["data"]["document_version"]["original_filename"] == real_name


def test_upload_with_mismatched_magic_bytes_is_rejected(api, db, owner):
    """34.16 — the declared content type is a claim, not a fact. This is the case
    the sniffing exists for."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    response = api.post(f"{V1}/contracts/{contract_id}/document-versions",
                        content=b"#!/bin/sh\nrm -rf /\n",
                        headers={"Content-Type": "application/pdf",
                                 "X-Filename": "not-really.pdf"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "UPLOAD_REJECTED"
    assert db.execute(select(M.DocumentVersion)).first() is None


def test_reupload_reports_the_duplicate_rather_than_suppressing_it(api, db, owner):
    """34.5 / Step 33.9 — whether a re-upload is a new contractual version is a
    business decision, so the duplicate is reported, never silently dropped."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    payload = build_docx(["1. Limitation of Liability"])
    headers = {"Content-Type": DOCX_MIME, "X-Filename": "msa.docx"}

    first = api.post(f"{V1}/contracts/{contract_id}/document-versions",
                     content=payload, headers=headers).json()["data"]
    second = api.post(f"{V1}/contracts/{contract_id}/document-versions",
                      content=payload, headers=headers).json()["data"]
    assert second["duplicate_of"] == first["document_version"]["id"]
    assert second["document_version"]["version_number"] == 2


def test_download_returns_the_preserved_original(api, db, owner):
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    payload = build_docx(["1. Limitation of Liability"])
    version_id = api.post(
        f"{V1}/contracts/{contract_id}/document-versions", content=payload,
        headers={"Content-Type": DOCX_MIME, "X-Filename": "msa.docx"}
    ).json()["data"]["document_version"]["id"]

    response = api.get(f"{V1}/document-versions/{version_id}/content")
    assert response.status_code == 200
    # 34.2 / 34.5 — byte-for-byte the original.
    assert response.content == payload
    assert response.headers["X-Content-Type-Options"] == "nosniff"
    assert "attachment" in response.headers["Content-Disposition"]


def test_download_is_a_permission_distinct_from_view(api, db, owner):
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = api.post(
        f"{V1}/contracts/{contract_id}/document-versions",
        content=build_docx(["1. Liability"]),
        headers={"Content-Type": DOCX_MIME, "X-Filename": "msa.docx"}
    ).json()["data"]["document_version"]["id"]

    # A second identity owning nothing cannot even see it (41.24)...
    stranger = make_user(db)
    grant_role(db, stranger, P.ROLE_USER)
    sign_in(api, db, stranger)
    assert api.get(f"{V1}/document-versions/{version_id}").status_code == 404

    # ...while the owner without document.download is refused the bytes but not
    # the metadata.
    narrow = bespoke_role(db, "VIEW_NOT_DOWNLOAD",
                          [P.CONTRACT_VIEW, P.DOCUMENT_VIEW])
    limited = make_user(db)
    grant(db, limited, narrow)
    contract = db.get(M.Contract, __import__("uuid").UUID(contract_id))
    contract.owner_id = limited.id
    db.flush()
    sign_in(api, db, limited)
    assert api.get(f"{V1}/document-versions/{version_id}").status_code == 200
    assert api.get(f"{V1}/document-versions/{version_id}/content"
                   ).status_code == 403


# =====================================================================
# Reviews — 49.8 idempotency
# =====================================================================
def test_review_creation_is_idempotent(api, db, owner):
    """49.8 — idempotent on ``(document_version_id, configuration_snapshot_id)``,
    so a retry cannot produce two Reviews of the same document against the same
    configuration (43.28)."""
    review = make_review_for(db, owner)
    sign_in(api, db, owner)
    body = {"document_version_id": str(review.document_version_id),
            "configuration_snapshot_id": str(review.configuration_snapshot_id)}

    first = api.post(f"{V1}/reviews", json=body)
    second = api.post(f"{V1}/reviews", json=body)
    assert first.json()["data"]["id"] == second.json()["data"]["id"]
    assert db.execute(
        select(M.Review).where(
            M.Review.document_version_id == review.document_version_id)
    ).scalars().all().__len__() == 1


def test_review_creation_for_someone_elses_document_is_404(api, db, owner):
    stranger = make_user(db)
    grant_role(db, stranger, P.ROLE_USER)
    theirs = make_review_for(db, stranger)

    sign_in(api, db, owner)
    response = api.post(f"{V1}/reviews", json={
        "document_version_id": str(theirs.document_version_id),
        "configuration_snapshot_id": str(theirs.configuration_snapshot_id)})
    assert response.status_code == 404


# =====================================================================
# Report — F-8 / F-9 / 36.10
# =====================================================================
def test_report_carries_no_risk_score_or_verdict(api, db, owner,
                                                 requirement_version):
    """36.10 forbids a risk score as the primary V1 legal output; F-8 makes risk
    a *configured* display mapping. None is configured, and ENG-09 says an absent
    configuration value fails closed rather than defaulting — so the field is
    absent, not guessed."""
    from tests.conftest import make_evaluation, make_finding

    review = make_review_for(db, owner)
    finding = make_finding(db, review, requirement_version)
    make_evaluation(db, finding, rule_outcome=E.RuleOutcome.APPROVAL_REQUIRED)

    sign_in(api, db, owner)
    body = api.get(f"{V1}/reviews/{review.id}/report").json()["data"]

    for forbidden in ("risk", "risk_score", "risk_level", "verdict", "overall"):
        assert forbidden not in body
    assert body["classification_counts"] == {"DEVIATION": 1}
    assert body["findings_requiring_decision"] == 1
    assert body["coverage"]["requirements_with_findings"] == 1
    # F-9 — a ratio over evaluated Requirements, carrying no legal meaning.
    assert body["alignment"] == {"requirements_evaluated": 1, "matched": 0,
                                "ratio": 0.0}


# =====================================================================
# Configuration — Step 29, 42.12, ENG-09
# =====================================================================
def _legal_admin(db):
    user = make_user(db)
    grant_role(db, user, P.ROLE_DEPARTMENT_LEAD)
    return user


def test_publish_fails_closed_when_configuration_is_incomplete(api, db, seeded):
    """ENG-09 — publishing a partial snapshot would produce Reviews that silently
    skipped a Requirement, so it is refused and the Requirement is named."""
    admin = _legal_admin(db)
    requirement = M.Requirement(code="LIABILITY-001", status=E.ConfigStatus.ACTIVE)
    db.add(requirement); db.flush()
    db.add(M.RequirementVersion(
        requirement_id=requirement.id, version_number=1, name="Liability",
        evaluator_type=E.EvaluatorType.NUMERIC_COMPARISON, created_by=admin.id))
    db.flush()

    sign_in(api, db, admin)
    response = api.post(f"{V1}/configuration/publish", json={})
    assert response.status_code == 422
    assert "LIABILITY-001" in response.json()["error"]["message"]
    assert db.execute(select(M.ConfigurationSnapshot)).first() is None


def test_draft_and_publish_produces_a_stable_snapshot(api, db, seeded):
    """Step 29 / rule 16 — a draft never affects a comparison until published,
    and republishing identical configuration reuses the identical snapshot, which
    is what keeps determinism (rule 9)."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)

    created = api.post(f"{V1}/requirements", json={"code": "LIABILITY-001"})
    assert created.status_code == 201
    requirement_id = created.json()["data"]["id"]

    # NOTE: the payloads below are STRUCTURAL placeholders supplied by the test,
    # not a legal position. Rule 21 — real Company Standards and Legal Rules must
    # come from the owner; this asserts the mechanism, not any legal content.
    versioned = api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Limitation of Liability",
        "evaluator_type": "NUMERIC_COMPARISON",
        # document_type is required configuration (Step 28 scoping, owner Q2/Q3
        # of 2026-08-19): publish refuses an untyped standard, so the happy path
        # must state it — same reasoning as confirm_threshold below.
        "company_standard": {"unit": "months", "document_type": "MSA"},
        # confirm_threshold is required configuration (D-1): publishing without
        # one is refused, so the happy path must state it.
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {"scope_required": True},
        "legal_rule": {"rule_type": "THRESHOLD", "configuration": {}},
    })
    assert versioned.status_code == 201

    # Still DRAFT, so it is not in a snapshot yet.
    assert api.post(f"{V1}/configuration/publish",
                    json={}).status_code == 422

    first = api.post(f"{V1}/configuration/publish",
                     json={"requirement_codes": ["LIABILITY-001"]})
    assert first.status_code == 201
    assert first.json()["data"]["requirement_count"] == 1
    assert first.json()["data"]["reused_existing"] is False

    again = api.post(f"{V1}/configuration/publish", json={})
    assert again.json()["data"]["id"] == first.json()["data"]["id"]
    assert again.json()["data"]["reused_existing"] is True


def test_publish_refuses_mapping_rules_without_a_confirm_threshold(api, db, seeded):
    """D-1 / ENG-09 — an unusable mapping rule version is as incomplete as an
    absent one.

    Refusing at publish keeps it out of every snapshot, so analysis never faces the
    case. The alternative — mapping as UNRESOLVED — produces UNABLE_TO_EVALUATE,
    which is Tier 1 and therefore *requires a Legal Decision* under D-3.5(b): a
    missing configuration number would create work in the Legal queue, clearable
    only by an authorized person ruling on a Requirement that was never evaluated.
    """
    admin = _legal_admin(db)
    sign_in(api, db, admin)

    requirement_id = api.post(f"{V1}/requirements",
                              json={"code": "LIABILITY-002"}).json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Limitation of Liability",
        "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {},
        "mapping_rules": {"exact_phrases": ["limitation of liability"]},
        "evaluation_rules": {},
    })

    response = api.post(f"{V1}/configuration/publish",
                        json={"requirement_codes": ["LIABILITY-002"]})
    assert response.status_code == 422
    body = response.json()["error"]
    assert "LIABILITY-002" in body["message"]
    assert "confirm_threshold" in body["message"]
    # Nothing was pinned: an incomplete Requirement cannot enter a snapshot.
    assert db.execute(select(M.ConfigurationSnapshot)).first() is None


def test_a_new_requirement_version_never_edits_the_previous_one(api, db, seeded):
    """Rule 16 / AUD-04 — a historical Review must reproduce from the exact
    versions it used, which is only possible if versions are appended."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    requirement_id = api.post(f"{V1}/requirements",
                              json={"code": "TERMINATION-001"}
                              ).json()["data"]["id"]
    draft = {"name": "Termination", "evaluator_type": "PRESENCE",
             "company_standard": {}, "mapping_rules": {},
             "evaluation_rules": {}}
    api.post(f"{V1}/requirements/{requirement_id}/versions", json=draft)
    api.post(f"{V1}/requirements/{requirement_id}/versions",
             json={**draft, "name": "Termination (revised)"})

    body = api.get(f"{V1}/requirements/{requirement_id}").json()["data"]
    assert [v["version_number"] for v in body["versions"]] == [1, 2]
    assert body["versions"][0]["name"] == "Termination"


def test_configuration_view_does_not_confer_draft(api, db, seeded):
    """Step 23 / 47.4 — Legal Reviewer may view configuration but not author it."""
    reviewer = make_user(db)
    grant_role(db, reviewer, P.ROLE_LEGAL_REVIEWER)
    sign_in(api, db, reviewer)
    assert api.get(f"{V1}/requirements").status_code == 200
    assert api.post(f"{V1}/requirements", json={"code": "X-1"}).status_code == 403


def test_legal_admin_cannot_publish_without_the_permission(api, db, seeded):
    narrow = make_user(db)
    grant(db, narrow, bespoke_role(db, "DRAFT_ONLY",
                                   [P.CONFIGURATION_VIEW, P.CONFIGURATION_DRAFT]))
    sign_in(api, db, narrow)
    assert api.post(f"{V1}/configuration/publish", json={}).status_code == 403


# =====================================================================
# Administration — S-8, S-9, S-10, SEC-05
# =====================================================================
def _super_admin(db):
    user = make_user(db)
    grant_role(db, user, P.ROLE_PLATFORM_ADMIN)
    return user


def test_a_new_account_holds_no_roles(api, db, seeded):
    """47.1.3 r3 — authority is always a later, deliberate grant."""
    sign_in(api, db, _super_admin(db))
    created = api.post(f"{V1}/users", json={"email": "New.User@Example.test",
                                            "name": "New User"})
    assert created.status_code == 201
    assert created.json()["data"]["roles"] == []
    assert created.json()["data"]["email"] == "new.user@example.test"


def test_cannot_grant_an_authority_one_does_not_hold(api, db, seeded):
    """S-8 — the vulnerability the external reference describes, in LegalMind's
    terms: any administrator able to grant the highest privilege through an
    ordinary role assignment defeats ROLE-05 entirely."""
    admin = _super_admin(db)
    target = make_user(db)
    sign_in(api, db, admin)

    response = api.post(f"{V1}/users/{target.id}/roles",
                        json={"role_code": P.ROLE_LEGAL_DECISION_AUTHORITY})
    assert response.status_code == 403
    assert "escalation refused" in response.json()["error"]["message"]


def test_granting_legal_authority_is_audited_as_such(api, db, seeded):
    """47.9 — a legal-authority change gets its own action, so it is findable
    without knowing which role codes happen to carry ``legal.*``."""
    granter = make_user(db)
    grant(db, granter, bespoke_role(
        db, "AUTHORITY_GRANTER",
        [P.USER_MANAGE, P.LEGAL_DECISION, P.LEGAL_APPROVE_CUSTOMIZATION]))
    target = make_user(db)

    sign_in(api, db, granter)
    response = api.post(f"{V1}/users/{target.id}/roles",
                        json={"role_code": P.ROLE_LEGAL_DECISION_AUTHORITY})
    assert response.status_code == 201
    assert P.ROLE_LEGAL_DECISION_AUTHORITY in response.json()["data"]["roles"]

    actions = set(db.execute(select(M.AuditEvent.action)).scalars().all())
    assert "admin.role_granted" in actions
    assert "admin.legal_authority_granted" in actions


def test_cannot_administer_a_more_privileged_account(api, db, seeded):
    """S-9 — the guard covers editing, not only granting. Without it an
    administrator who cannot grant ``legal.decision`` could disable the account
    that holds it and reach the same outcome."""
    admin = _super_admin(db)
    authority = make_user(db)
    grant(db, authority, bespoke_role(db, "AUTH_HOLDER", [P.LEGAL_DECISION]))

    sign_in(api, db, admin)
    response = api.patch(f"{V1}/users/{authority.id}",
                         json={"status": "DISABLED"})
    assert response.status_code == 403


def test_disabling_the_last_legal_authority_is_refused(api, db, seeded):
    """SEC-05 — never zero legal authorities.

    Only ACTIVE users count: a disabled account cannot authenticate by any route
    (47.1.3), so counting its grant would satisfy SEC-05 on paper while leaving
    exactly the stalled Review the rule exists to prevent (Step 31 r18).
    """
    holder = make_user(db)
    grant(db, holder, bespoke_role(db, "SOLE_AUTHORITY",
                                   [P.LEGAL_DECISION, P.USER_MANAGE]))
    sign_in(api, db, holder)
    response = api.patch(f"{V1}/users/{holder.id}", json={"status": "DISABLED"})
    assert response.status_code == 403
    assert "Legal" in response.json()["error"]["message"]
    assert db.get(M.User, holder.id).status is E.UserStatus.ACTIVE


def test_revoking_the_last_legal_authority_role_is_refused(api, db, seeded):
    holder = make_user(db)
    role = bespoke_role(db, "SOLE_AUTHORITY_2",
                        [P.LEGAL_DECISION, P.USER_MANAGE])
    grant(db, holder, role)
    sign_in(api, db, holder)
    response = api.delete(f"{V1}/users/{holder.id}/roles/{role.code}")
    assert response.status_code == 403


def test_disabling_an_account_revokes_its_sessions_immediately(api, db, seeded):
    """S-2 — a disabled account must not keep working for the rest of a live
    session."""
    admin = _super_admin(db)
    target = make_user(db)
    grant_role(db, target, P.ROLE_USER)

    from legalmind.security.sessions import create_session
    victim_session = create_session(db, target)

    sign_in(api, db, admin)
    assert api.patch(f"{V1}/users/{target.id}",
                     json={"status": "DISABLED"}).status_code == 200
    db.refresh(victim_session)
    assert victim_session.revoked_at is not None


def test_role_permission_replacement_refuses_escalation(api, db, seeded):
    """S-8 applied to permissions rather than roles — otherwise PATCH /roles
    would be the route that bypasses ``require_can_grant_role``."""
    admin = _super_admin(db)
    role = bespoke_role(db, "ORDINARY", [P.CONTRACT_VIEW])
    sign_in(api, db, admin)

    response = api.patch(f"{V1}/roles/{role.id}",
                         json={"permissions": [P.CONTRACT_VIEW,
                                               P.LEGAL_DECISION]})
    assert response.status_code == 403
    assert set(db.execute(
        select(M.Permission.name)
        .join(M.RolePermission,
              M.RolePermission.permission_id == M.Permission.id)
        .where(M.RolePermission.role_id == role.id)
    ).scalars().all()) == {P.CONTRACT_VIEW}


def test_role_permission_replacement_is_audited_with_before_and_after(api, db,
                                                                     seeded):
    """S-10 / 43.26 — the whole set changes or none of it does, and the change is
    recorded on both sides so an authority change is reconstructable."""
    admin = _super_admin(db)
    role = bespoke_role(db, "ORDINARY_2", [P.CONTRACT_VIEW])
    sign_in(api, db, admin)

    response = api.patch(f"{V1}/roles/{role.id}",
                         json={"permissions": [P.AUDIT_VIEW, P.USER_MANAGE]})
    assert response.status_code == 200
    assert response.json()["data"]["permissions"] == sorted(
        [P.AUDIT_VIEW, P.USER_MANAGE])

    event = db.execute(
        select(M.AuditEvent)
        .where(M.AuditEvent.action == "admin.permission_changed",
               M.AuditEvent.entity_id == role.id)
        .order_by(M.AuditEvent.id)
    ).scalars().all()[-1]
    assert event.before_state == {"permissions": [P.CONTRACT_VIEW]}
    assert P.USER_MANAGE in event.after_state["permissions"]


def test_unknown_permission_name_is_rejected(api, db, seeded):
    admin = _super_admin(db)
    role = bespoke_role(db, "ORDINARY_3", [])
    sign_in(api, db, admin)
    response = api.patch(f"{V1}/roles/{role.id}",
                         json={"permissions": ["contract.obliterate"]})
    assert response.status_code == 422


def test_role_listing_marks_which_roles_confer_legal_authority(api, db, seeded):
    """Makes the ROLE-05 boundary visible to an administrator without them having
    to know which permission names are special."""
    admin = make_user(db)
    grant(db, admin, bespoke_role(db, "ROLE_ADMIN_ONLY", [P.ROLE_MANAGE]))
    sign_in(api, db, admin)

    roles = {r["code"]: r for r in
             api.get(f"{V1}/roles", params={"page_size": 100}).json()["data"]}
    assert roles[P.ROLE_LEGAL_DECISION_AUTHORITY]["confers_legal_authority"] == [
        P.LEGAL_APPROVE_CUSTOMIZATION, P.LEGAL_DECISION]
    # Locked Step 23 — Super Admin confers none of it.
    assert roles[P.ROLE_PLATFORM_ADMIN]["confers_legal_authority"] == []


# =====================================================================
# S-5 / 49.10 — rate limiting
# =====================================================================
def test_login_is_rate_limited(api, db, seeded):
    """49.10 — applied to authentication. Thresholds are deployment
    configuration, so the test drives the configured limit rather than asserting
    a specific number."""
    from legalmind.api import ratelimit
    from legalmind.api.routers import auth as auth_router

    for _ in range(ratelimit.LOGIN.max_requests):
        assert api.post(f"{V1}/auth/login",
                        json={"email": "nobody@example.test",
                              "password": "x"}).status_code == 401

    limited = api.post(f"{V1}/auth/login",
                       json={"email": "nobody@example.test", "password": "x"})
    assert limited.status_code == 429
    assert limited.json()["error"]["code"] == "RATE_LIMITED"
    # 49.10 — no detail about the limit's shape.
    assert "Retry-After" not in limited.headers
    assert "limit" not in limited.json()["error"]["message"].lower()
    assert isinstance(auth_router.limiter, ratelimit.InProcessRateLimiter)


# ----------------------------------------------- Document Type at the boundaries
def test_publish_refuses_a_standard_without_a_document_type(api, db, seeded):
    """Step 28 scoping (owner Q2/Q3, 2026-08-19). Refusing at publish is what
    lets the analysis filter be a plain equality test: every snapshot item is
    guaranteed a valid type."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "UNTYPED-001"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Untyped", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"unit": "months"},          # no document_type
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {},
    })
    response = api.post(f"{V1}/configuration/publish",
                        json={"requirement_codes": ["UNTYPED-001"]})
    assert response.status_code == 422
    assert "document_type" in response.json()["error"]["message"]


def test_publish_refuses_an_unknown_document_type(api, db, seeded):
    """Locked Step 6 fixes the vocabulary; a near-miss is refused, not coerced."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "MISTYPED-001"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Mistyped", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"unit": "months", "document_type": "msa"},  # lowercase
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {},
    })
    response = api.post(f"{V1}/configuration/publish",
                        json={"requirement_codes": ["MISTYPED-001"]})
    assert response.status_code == 422
    assert "msa" in response.json()["error"]["message"]


def test_contract_creation_refuses_a_type_outside_step_6(api, db, owner):
    """The uploader declares the type (owner Q9) from Step 6's ten values."""
    sign_in(api, db, owner)
    bad = api.post(f"{V1}/contracts",
                   json={"name": "X", "contract_type": "LEASE"})
    assert bad.status_code == 422
    good = api.post(f"{V1}/contracts",
                    json={"name": "X", "contract_type": "NDA"})
    assert good.status_code == 201
    assert good.json()["data"]["contract_type"] == "NDA"


# ------------------------------------------- admin standards surface (2026-08-19)
def test_requirement_detail_returns_values_and_authorship(api, db, seeded):
    """The admin read path: an admin screen must be able to display "current:
    12 months, changed by X". Values were previously write-only through the API,
    which made the stored configuration unreviewable."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "READBACK-001"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Readback", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"document_type": "MSA", "preferred": 6,
                             "unit": "MONTHS"},
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {},
    })
    detail = api.get(f"{V1}/requirements/{requirement_id}").json()["data"]
    version = detail["versions"][0]
    assert version["company_standard"]["preferred"] == 6
    assert version["company_standard"]["document_type"] == "MSA"
    assert version["created_by"] == str(admin.id)
    # The values-free list view stays values-free.
    listed = api.get(f"{V1}/requirements").json()["data"]
    assert all("company_standard" not in v
               for r in listed for v in r["versions"])


def test_standard_update_appends_and_never_edits(api, db, seeded):
    """Locked rule 16 through the new admin path: "edit and save" appends a new
    version carrying the other artifacts forward; the prior version's values
    stay byte-identical, so historical Reviews remain reproducible."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "EDIT-001"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "Edit target", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"document_type": "MSA", "preferred": 6,
                             "unit": "MONTHS"},
        "mapping_rules": {"exact_phrase": ["x"], "confirm_threshold": 5},
        "evaluation_rules": {"scope_required": True},
    })

    updated = api.post(f"{V1}/requirements/{requirement_id}/standard", json={
        "company_standard": {"document_type": "MSA", "preferred": 18,
                             "unit": "MONTHS"},
        "reason": "structural test: exercise the append path",
    })
    assert updated.status_code == 201
    versions = updated.json()["data"]["versions"]
    assert [v["version_number"] for v in versions] == [1, 2]
    assert versions[0]["company_standard"]["preferred"] == 6   # untouched
    assert versions[1]["company_standard"]["preferred"] == 18

    # Mapping rules were carried forward unchanged, not re-supplied.
    rv2 = db.execute(
        select(M.RequirementVersion)
        .where(M.RequirementVersion.version_number == 2)
    ).scalars().first()
    mr2 = db.execute(
        select(M.MappingRuleVersion)
        .where(M.MappingRuleVersion.requirement_version_id == rv2.id)
    ).scalars().one()
    assert mr2.rules == {"exact_phrase": ["x"], "confirm_threshold": 5}


def test_standard_update_refuses_untyped_and_requires_reason(api, db, seeded):
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "EDIT-002"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "T", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"document_type": "MSA", "unit": "MONTHS"},
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {},
    })
    # No document_type in the replacement → refused at save, not at publish.
    r = api.post(f"{V1}/requirements/{requirement_id}/standard", json={
        "company_standard": {"unit": "MONTHS"}, "reason": "x"})
    assert r.status_code == 422
    # No reason → schema refusal.
    r = api.post(f"{V1}/requirements/{requirement_id}/standard", json={
        "company_standard": {"document_type": "MSA", "unit": "MONTHS"}})
    assert r.status_code == 422


def test_configuration_writes_are_audited(api, db, seeded):
    """Closing the gap where every other privileged router recorded audit events
    and configuration writes did not. A standard change is a change of legal
    position; "who changed what, when, why" must be answerable from the trail."""
    admin = _legal_admin(db)
    sign_in(api, db, admin)
    created = api.post(f"{V1}/requirements", json={"code": "AUDITED-001"})
    requirement_id = created.json()["data"]["id"]
    api.post(f"{V1}/requirements/{requirement_id}/versions", json={
        "name": "A", "evaluator_type": "NUMERIC_COMPARISON",
        "company_standard": {"document_type": "MSA", "unit": "MONTHS"},
        "mapping_rules": {"exact_phrase": [], "confirm_threshold": 5},
        "evaluation_rules": {},
    })
    api.post(f"{V1}/requirements/{requirement_id}/standard", json={
        "company_standard": {"document_type": "MSA", "preferred": 9,
                             "unit": "MONTHS"},
        "reason": "structural audit-trail test"})
    api.post(f"{V1}/configuration/publish",
             json={"requirement_codes": ["AUDITED-001"]})

    actions = [row.action for row in db.execute(
        select(M.AuditEvent).order_by(M.AuditEvent.timestamp)).scalars()]
    for expected in ("config.requirement_created", "config.version_created",
                     "config.standard_updated", "config.published"):
        assert expected in actions, f"missing audit action {expected}"

    updated = next(row for row in db.execute(
        select(M.AuditEvent)
        .where(M.AuditEvent.action == "config.standard_updated")).scalars())
    assert updated.after_state["reason"] == "structural audit-trail test"
    assert updated.actor_id == admin.id
    # 53.3 — the audit event names the change, never the values.
    assert "preferred" not in json.dumps(updated.after_state)


def test_import_tool_writes_ratified_standards_idempotently(db, seeded):
    """tools/import_ratified_standards — the bridge from ratified config files
    to the runtime database. Idempotent by content: run twice, one version."""
    from tests.conftest import make_user
    from tools.import_ratified_standards import import_standards

    actor = make_user(db)
    first = import_standards(db, actor_email=actor.email)
    assert any("LIABILITY-MSA-001: version 1 written" in line for line in first)
    assert any("LIABILITY-TOS-001: version 1 written" in line for line in first)
    # Since 2026-08-19 every ratified file carries mapping and evaluation rules
    # (the owner's publishability tasking), so nothing imports as an
    # unpublishable draft any more. The refusal path stays covered: a file
    # without rules would still be reported, and the publish gate re-checks.
    assert not any("unpublishable draft" in line for line in first)

    second = import_standards(db, actor_email=actor.email)
    assert any("LIABILITY-MSA-001: unchanged" in line for line in second)

    msa = db.execute(
        select(M.Requirement).where(M.Requirement.code == "LIABILITY-MSA-001")
    ).scalars().one()
    versions = db.execute(
        select(M.RequirementVersion)
        .where(M.RequirementVersion.requirement_id == msa.id)).scalars().all()
    assert len(versions) == 1                        # idempotent
    cs = db.execute(
        select(M.CompanyStandardVersion)
        .where(M.CompanyStandardVersion.requirement_version_id == versions[0].id)
    ).scalars().one()
    assert cs.configuration["preferred"] == 12       # Constitution §9 (AM-43 r4, 2026-09-08)
    assert cs.configuration["document_type"] == "MSA"
    assert ("Legal Constitution" in versions[0].description   # AM-43 r4: the MSA standard is sourced from the Constitution since 2026-09-08
            or "Master Services" in versions[0].description)

    # The terminology of 2026-08-19 makes the Requirement publishable: mapping
    # rules with a usable confirm_threshold (D-1) and an evaluation rule version.
    mr = db.execute(
        select(M.MappingRuleVersion)
        .where(M.MappingRuleVersion.requirement_version_id == versions[0].id)
    ).scalars().one()
    assert isinstance(mr.rules.get("confirm_threshold"), int)
    er = db.execute(
        select(M.EvaluationRuleVersion)
        .where(M.EvaluationRuleVersion.requirement_version_id == versions[0].id)
    ).scalars().one()
    assert er.rules

    # The approved zero-tolerance Legal Rule (owner, 2026-08-20) is written from
    # the file — verbatim, and nothing else would have been accepted.
    lr = db.execute(
        select(M.LegalRuleVersion)
        .where(M.LegalRuleVersion.requirement_version_id == versions[0].id)
    ).scalars().one()
    assert lr.configuration == {"deviation_outcome": "UNACCEPTABLE",
                                "unlimited_outcome": "UNACCEPTABLE"}


def test_every_ratified_standard_publishes_through_the_gated_endpoint(api, db, seeded):
    """The tasked outcome, proved end to end rather than asserted from the files.

    The owner's tasking of 2026-08-19 was to make the ratified Requirements
    publishable. This imports all of them and publishes through the real
    permission-gated endpoint, so the publish gate's own fail-closed checks —
    company standard, mapping rules with a usable `confirm_threshold` (D-1),
    evaluation rules, and a locked Step 6 `document_type` — are what pass, not a
    re-implementation of them in a test.

    Rule 21 note: nothing here authors legal content. The values come from the
    owner's ratified files; the assertion is about the mechanism reaching them.
    """
    from legalmind.evaluation.corpus import RATIFIED_STANDARDS_DIR
    from tools.import_ratified_standards import import_standards

    admin = _legal_admin(db)
    sign_in(api, db, admin)

    import_standards(db, actor_email=admin.email)
    db.flush()

    codes = sorted(p.stem for p in RATIFIED_STANDARDS_DIR.glob("*.json"))
    published = api.post(f"{V1}/configuration/publish",
                         json={"requirement_codes": codes})
    # A refusal names the incomplete Requirement, so surface it rather than a
    # bare status code.
    assert published.status_code == 201, published.text
    assert published.json()["data"]["requirement_count"] == len(codes)


# ---------------------------------------------------------------------------
# Contract detail carries its document versions (implementation addition, 2026-08-30)
# ---------------------------------------------------------------------------
def test_contract_detail_lists_its_document_versions_newest_first(api, db, owner):
    """A document-anchored workspace opened on a contract must find its document
    through the API. Newest first, and the shape is the existing version
    serializer's — nothing new is disclosed and `storage_key` stays absent."""
    sign_in(api, db, owner)
    contract = api.post("/api/v1/contracts",
                        json={"name": "Versioned", "contract_type": "MSA"}).json()["data"]
    # The field lives on the DETAIL endpoint (like `assist_index` on a version):
    # the list endpoint stays lean, and a create returns the contract it made.
    empty = api.get(f"/api/v1/contracts/{contract['id']}").json()["data"]
    assert empty["document_versions"] == []

    docx = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    for n in (1, 2):
        response = api.post(
            f"/api/v1/contracts/{contract['id']}/document-versions",
            content=build_docx([f"Clause {n}. Liability is limited to fees paid."]),
            headers={"content-type": docx, "x-filename": f"msa-v{n}.docx"})
        assert response.status_code in (200, 201), response.text

    detail = api.get(f"/api/v1/contracts/{contract['id']}").json()["data"]
    versions = detail["document_versions"]
    assert [v["version_number"] for v in versions] == [2, 1]
    assert versions[0]["original_filename"] == "msa-v2.docx"
    assert "storage_key" not in versions[0]
    assert "processing_status" in versions[0]


# =====================================================================
# 2026-08-31 UX correction — the two additive reads (Step 49 record)
# =====================================================================
def test_snapshots_list_is_metadata_only_and_needs_review_create(api, db, owner):
    """`GET /configuration/snapshots` exists so a client can resolve "analyze
    against the current standards" to a concrete snapshot id. Newest first,
    metadata ONLY — no snapshot item, no standard value (LEGAL-02)."""
    review = make_review_for(db, owner)   # creates one snapshot
    sign_in(api, db, owner)
    body = api.get(f"{V1}/configuration/snapshots").json()
    assert body["pagination"]["total"] >= 1
    newest = body["data"][0]
    assert set(newest) == {"id", "snapshot_hash", "created_at", "requirement_count"}
    assert str(review.configuration_snapshot_id) in [r["id"] for r in body["data"]]

    # A caller without review.create has no business listing them.
    bare = make_user(db)
    sign_in(api, db, bare)
    assert api.get(f"{V1}/configuration/snapshots").status_code == 403


def test_documents_list_carries_a_permission_layered_analysis_summary(
        api, db, owner, requirement_version):
    """2026-08-31 UX correction: each list row answers "what did analysis find"
    — the latest version's latest Review with classification counts — instead
    of echoing a lifecycle enum. Counts need `finding.view` and are OMITTED
    (never nulled) without it — Step 24 r8 applied to a projection."""
    from tests.conftest import make_finding

    review = make_review_for(db, owner)
    make_finding(db, review, requirement_version,
                 classification=E.FindingClassification.DEVIATION)
    sign_in(api, db, owner)
    rows = api.get(f"{V1}/contracts").json()["data"]
    row = next(r for r in rows if r["id"] == str(review.contract_id))
    assert row["latest_version"]["version_number"] == 1
    assert row["latest_version"]["processing_status"] == "COMPLETED"
    analysis = row["latest_analysis"]
    assert analysis["review_id"] == str(review.id)
    assert analysis["review_status"] == "ANALYSIS_COMPLETE"
    assert analysis["classification_counts"] == {"DEVIATION": 1}

    # A contract with no Review states that honestly rather than inventing one.
    bare = api.post(f"{V1}/contracts", json={"name": "Fresh", "contract_type": "MSA"})
    fresh = next(r for r in api.get(f"{V1}/contracts").json()["data"]
                 if r["id"] == bare.json()["data"]["id"])
    assert fresh["latest_version"] is None
    assert fresh["latest_analysis"] is None


# ==========================================================================
# Documents list — search / type filter / status filter / sort / summary
# (2026-09-01, owner-directed Documents redesign)
# ==========================================================================
def test_contracts_list_supports_name_search_and_type_filter(api, db, owner):
    sign_in(api, db, owner)
    api.post(f"{V1}/contracts", json={"name": "Leapswitch GRP MSA", "contract_type": "MSA"})
    api.post(f"{V1}/contracts", json={"name": "CloudPe Terms of Service", "contract_type": "TOS"})

    by_name = api.get(f"{V1}/contracts", params={"q": "leapswitch"}).json()["data"]
    assert [r["name"] for r in by_name] == ["Leapswitch GRP MSA"]

    by_type = api.get(f"{V1}/contracts", params={"contract_type": "TOS"}).json()["data"]
    assert [r["name"] for r in by_type] == ["CloudPe Terms of Service"]

    assert api.get(f"{V1}/contracts", params={"q": "no such contract"}).json()["data"] == []


def test_contracts_list_sorts_by_name(api, db, owner):
    sign_in(api, db, owner)
    api.post(f"{V1}/contracts", json={"name": "Zeta MSA", "contract_type": "MSA"})
    api.post(f"{V1}/contracts", json={"name": "Alpha MSA", "contract_type": "MSA"})
    rows = api.get(f"{V1}/contracts", params={"sort": "name_asc"}).json()["data"]
    names = [r["name"] for r in rows]
    assert names.index("Alpha MSA") < names.index("Zeta MSA")


def test_contracts_list_status_filter_matches_the_same_bucket_the_row_shows(
        api, db, owner, requirement_version):
    """The bucket a row would compute for itself (draft/analyzing/needs_attention/
    analyzed) is exactly what `?status=` filters on — one derivation, reused,
    never a second one that could disagree (52.7's rule applied to filtering)."""
    from tests.conftest import make_finding

    review = make_review_for(db, owner)  # ANALYSIS_COMPLETE, no findings yet -> analyzed
    sign_in(api, db, owner)
    api.post(f"{V1}/contracts", json={"name": "Draft Only", "contract_type": "MSA"})

    analyzed = api.get(f"{V1}/contracts", params={"status": "analyzed"}).json()["data"]
    assert {r["id"] for r in analyzed} == {str(review.contract_id)}

    draft = api.get(f"{V1}/contracts", params={"status": "draft"}).json()["data"]
    assert "Draft Only" in {r["name"] for r in draft}
    assert str(review.contract_id) not in {r["id"] for r in draft}

    # A DEVIATION finding moves the SAME contract from analyzed -> needs_attention.
    make_finding(db, review, requirement_version,
                classification=E.FindingClassification.DEVIATION)
    db.flush()
    needs_attention = api.get(f"{V1}/contracts", params={"status": "needs_attention"}).json()["data"]
    assert {r["id"] for r in needs_attention} == {str(review.contract_id)}
    assert api.get(f"{V1}/contracts", params={"status": "analyzed"}).json()["data"] == []

    assert api.get(f"{V1}/contracts", params={"status": "not-a-real-bucket"}).status_code == 422 \
        or api.get(f"{V1}/contracts", params={"status": "not-a-real-bucket"}).status_code == 400


def test_contracts_summary_counts_match_the_list_buckets_across_every_contract(
        api, db, owner, requirement_version):
    """The summary tiles and the list's own `?status=` filter must never
    disagree — both reduce to the identical `_status_bucket` computation."""
    from tests.conftest import make_finding

    review = make_review_for(db, owner)
    make_finding(db, review, requirement_version,
                classification=E.FindingClassification.MISSING)
    sign_in(api, db, owner)
    api.post(f"{V1}/contracts", json={"name": "Untouched Draft", "contract_type": "MSA"})

    summary = api.get(f"{V1}/contracts/summary").json()["data"]
    assert summary["total"] == 2
    assert summary["needs_attention"] == 1
    assert summary["draft"] == 1
    assert summary["analyzing"] == 0
    assert summary["analyzed"] == 0

    for bucket in ("draft", "analyzing", "needs_attention", "analyzed"):
        listed = api.get(f"{V1}/contracts", params={"status": bucket}).json()["data"]
        assert len(listed) == summary[bucket]


def test_contracts_summary_requires_contract_view(api, db):
    from tests.conftest import make_user

    bare = make_user(db)
    sign_in(api, db, bare)
    assert api.get(f"{V1}/contracts/summary").status_code == 403


# =====================================================================
# Declared version metadata — source · counterparty · effective date (2026-09-06)
# Locked 42.4 `metadata` JSONB; Step 2 "should store"; Step 6 source axis.
# =====================================================================
def _upload(api, contract_id):
    return api.post(
        f"{V1}/contracts/{contract_id}/document-versions",
        content=build_docx(["1. Limitation of Liability",
                            "Liability is capped at fees paid."]),
        headers={"Content-Type": DOCX_MIME, "X-Filename": "msa.docx"},
    ).json()["data"]["document_version"]["id"]


def test_declared_metadata_is_omitted_until_declared_then_round_trips(api, db, owner):
    """Absent, never null (the `SEC-07` discipline): a version nobody has
    described carries none of the three keys. Declared, they come back on the
    version AND on the Dashboard list's `latest_version`, so the counterparty
    datalist can converge on names the caller already sees."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = _upload(api, contract_id)

    before = api.get(f"{V1}/document-versions/{version_id}").json()["data"]
    assert not {"source", "counterparty", "effective_date"} & before.keys()

    declared = api.patch(f"{V1}/document-versions/{version_id}", json={
        "source": "COUNTERPARTY",
        "counterparty": "  Placeholder Counterparty Ltd ",   # trimmed, never a real name (54.6)
        "effective_date": "2026-07-28",
    })
    assert declared.status_code == 200, declared.text
    body = declared.json()["data"]
    assert body["source"] == "COUNTERPARTY"
    assert body["counterparty"] == "Placeholder Counterparty Ltd"
    assert body["effective_date"] == "2026-07-28"

    # Persisted as the locked JSONB, beside the key ingestion already writes.
    stored = db.get(M.DocumentVersion, uuid.UUID(version_id)).doc_metadata
    assert stored["source"] == "COUNTERPARTY" and stored["effective_date"] == "2026-07-28"

    listed = api.get(f"{V1}/contracts").json()["data"]
    row = next(c for c in listed if c["id"] == contract_id)
    assert row["latest_version"]["counterparty"] == "Placeholder Counterparty Ltd"
    assert row["latest_version"]["source"] == "COUNTERPARTY"

    # A field left out is untouched; a field sent as null is cleared AND omitted.
    cleared = api.patch(f"{V1}/document-versions/{version_id}",
                        json={"counterparty": None}).json()["data"]
    assert "counterparty" not in cleared
    assert cleared["source"] == "COUNTERPARTY"          # untouched


def test_declared_metadata_refuses_a_bad_source_and_a_bad_date(api, db, owner):
    """Refused at the boundary, not coerced — the same rule as Document Type:
    `counterparty` is not `COUNTERPARTY`, and the 40th of the 13th is not a date."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = _upload(api, contract_id)

    assert api.patch(f"{V1}/document-versions/{version_id}",
                     json={"source": "counterparty"}).status_code == 422
    assert api.patch(f"{V1}/document-versions/{version_id}",
                     json={"effective_date": "2026-13-40"}).status_code == 422
    assert api.patch(f"{V1}/document-versions/{version_id}",
                     json={"nonsense": 1}).status_code == 422       # extra="forbid"
    # Whitespace-only is "nothing declared", not a counterparty called " ".
    body = api.patch(f"{V1}/document-versions/{version_id}",
                     json={"counterparty": "   "}).json()["data"]
    assert "counterparty" not in body


def test_declared_metadata_is_fixed_once_a_review_exists(api, db, owner):
    """Owner ruling 2026-09-06 on locked 33.7 / 34.15 r3: declared metadata may
    be corrected only while the version has NO Review. The existence of a
    Review is the test — not its status, not its outcome — so a DRAFT Review
    already freezes it. 409, because the caller is the owner and existence is
    not the secret; the state is the reason."""
    sign_in(api, db, owner)
    review = make_review_for(db, owner)
    refused = api.patch(f"{V1}/document-versions/{review.document_version_id}",
                        json={"source": "ORGANIZATION"})
    assert refused.status_code == 409, refused.text
    assert db.get(M.DocumentVersion, review.document_version_id).doc_metadata is None


def test_declared_metadata_never_reaches_the_assist_lane():
    """`AM-30` t4 — no counterparty or organizational identifier may egress.
    The declared keys live only in `document_versions.metadata`; the assist
    package must not read that column, so a future prompt cannot pick a
    counterparty name up by accident. A static check, because the property is
    about what the code CAN do, not what one prompt happened to contain."""
    import pathlib

    assist = pathlib.Path(__file__).resolve().parents[1] / "legalmind" / "assist"
    offenders = [p.name for p in assist.glob("*.py")
                 if "doc_metadata" in p.read_text() or "declared_metadata" in p.read_text()]
    assert offenders == [], f"assist lane reads declared version metadata: {offenders}"


# =====================================================================
# P-8 — the document pane reads ONE processing run (2026-09-06)
# =====================================================================
def test_the_document_pane_reads_only_the_latest_completed_run(api, db, owner):
    """A retry or a REPROCESS writes a second segmentation of the same version;
    the pane must show ONE of them, never both merged. The older rows stay
    (42.5 history — a Finding may cite them); they are just not the document
    any more. And if the newest attempt FAILED, the last COMPLETED reading is
    still the document."""
    from datetime import UTC, datetime, timedelta

    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = uuid.UUID(_upload(api, contract_id))
    before = api.get(f"{V1}/document-versions/{version_id}/evidence").json()
    original_ids = {row["id"] for row in before["data"]}
    assert original_ids

    later = M.DocumentProcessingRun(
        document_version_id=version_id, run_type=E.ProcessingRunType.REPROCESS,
        status=E.ProcessingRunStatus.COMPLETED, processor_version="t",
        started_at=datetime.now(UTC) + timedelta(minutes=5))
    db.add(later); db.flush()
    db.add(M.DocumentEvidence(
        document_version_id=version_id, processing_run_id=later.id, page_number=1,
        section_number="1", section_title="Re-read", content="Re-segmented clause.",
        source_type=E.EvidenceSourceType.NATIVE_TEXT, start_offset=0, end_offset=20))
    db.flush()

    after = api.get(f"{V1}/document-versions/{version_id}/evidence").json()
    assert after["pagination"]["total"] == 1
    assert [row["content"] for row in after["data"]] == ["Re-segmented clause."]
    assert db.execute(
        select(func.count(M.DocumentEvidence.id))
        .where(M.DocumentEvidence.document_version_id == version_id)
    ).scalar_one() == len(original_ids) + 1          # history kept

    later.status = E.ProcessingRunStatus.FAILED
    db.flush()
    again = api.get(f"{V1}/document-versions/{version_id}/evidence").json()
    assert {row["id"] for row in again["data"]} == original_ids


# =====================================================================
# Phase 5 — re-read in place, Option C (owner, 2026-09-06)
# =====================================================================
def test_reprocess_rereads_with_the_current_parser_and_keeps_history(api, db, owner):
    """A NEW REPROCESS run over the preserved original; the earlier run's rows
    stay (42.5, rule 17) but the pane now shows the new reading (P-8); the
    assist index is rebuilt over the new rows; the act is audited."""
    from legalmind.assist import store as assist_store

    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = uuid.UUID(_upload(api, contract_id))
    before = api.get(f"{V1}/document-versions/{version_id}/evidence").json()["data"]
    old_ids = {row["id"] for row in before}
    chunks_before = assist_store.count_chunks(db, version_id)

    response = api.post(f"{V1}/document-versions/{version_id}/reprocess")
    assert response.status_code == 201, response.text
    body = response.json()["data"]
    assert body["processing_run"]["run_type"] == "REPROCESS"
    assert body["processing_run"]["status"] == "COMPLETED"
    assert body["evidence_count"] == len(old_ids)          # same file, same reading

    after = api.get(f"{V1}/document-versions/{version_id}/evidence").json()["data"]
    new_ids = {row["id"] for row in after}
    assert new_ids and not new_ids & old_ids                # the NEW run's rows
    runs = db.execute(select(M.DocumentProcessingRun)
                      .where(M.DocumentProcessingRun.document_version_id == version_id)
                      ).scalars().all()
    assert {r.run_type.value for r in runs} == {"PARSE", "REPROCESS"}
    assert db.execute(select(func.count(M.DocumentEvidence.id))
                      .where(M.DocumentEvidence.document_version_id == version_id)
                      ).scalar_one() == 2 * len(old_ids)  # history kept
    rerun = next(r for r in runs if r.run_type is E.ProcessingRunType.REPROCESS)
    assert rerun.run_metadata["reprocess_of"] == str(
        next(r.id for r in runs if r.run_type is E.ProcessingRunType.PARSE))
    assert assist_store.count_chunks(db, version_id) == chunks_before   # rebuilt, not doubled
    assert db.execute(select(M.AuditEvent)
                      .where(M.AuditEvent.action == "document.reprocessed")
                      ).scalars().one().entity_id == version_id


def test_reprocess_is_refused_while_anything_relies_on_the_current_reading(api, db, owner):
    """Option C's whole point. A Review, Key Obligations anchored to evidence, or a
    version still processing → 409 naming the reason and the way forward (a new
    version by re-upload). Nothing is written when refused."""
    from sqlalchemy import text

    from legalmind import config

    sign_in(api, db, owner)
    review = make_review_for(db, owner)
    refused = api.post(f"{V1}/document-versions/{review.document_version_id}/reprocess")
    assert refused.status_code == 409 and "Review" in refused.text
    assert "new version" in refused.text

    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = uuid.UUID(_upload(api, contract_id))
    evidence_id = db.execute(select(M.DocumentEvidence.id)
                             .where(M.DocumentEvidence.document_version_id == version_id)
                             .limit(1)).scalar_one()
    schema = config.assist_schema()
    run_id = uuid.uuid4()
    db.execute(text(f'INSERT INTO "{schema}".obligation_extraction_runs '
                    '(id, document_version_id, status, model_identity, prompt_version) '
                    "VALUES (:i, :d, 'COMPLETED', 'test', 'v')"),
               {"i": run_id, "d": version_id})
    db.execute(text(f'INSERT INTO "{schema}".obligation_extractions '
                    '(id, run_id, document_version_id, evidence_id, party_label, '
                    'obligation_text, ordinal) VALUES (:i, :r, :d, :e, :p, :t, 1)'),
               {"i": uuid.uuid4(), "r": run_id, "d": version_id, "e": evidence_id,
                "p": "Customer", "t": "shall pay"})
    db.flush()
    refused = api.post(f"{V1}/document-versions/{version_id}/reprocess")
    assert refused.status_code == 409 and "Key Obligations" in refused.text

    version = db.get(M.DocumentVersion, version_id)
    version.processing_status = E.ProcessingStatus.PROCESSING
    db.flush()
    assert api.post(f"{V1}/document-versions/{version_id}/reprocess").status_code == 409
    assert db.execute(select(func.count(M.DocumentProcessingRun.id))
                      .where(M.DocumentProcessingRun.document_version_id == version_id)
                      ).scalar_one() == 1                    # nothing written


# =====================================================================
# P-1 — the contract's lifecycle state is declared, and trailed (2026-09-06)
# =====================================================================
def test_contract_status_is_declared_by_the_owner_and_audited(api, db, owner):
    """Step 2's Draft / Active / Superseded through the PATCH that already
    existed. Any transition is allowed — a wrong click must be correctable — and
    each real change is one audit event; a no-op change writes none."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]

    assert api.patch(f"{V1}/contracts/{contract_id}",
                     json={"status": "ACTIVE"}).json()["data"]["status"] == "ACTIVE"
    assert api.patch(f"{V1}/contracts/{contract_id}",
                     json={"status": "ACTIVE"}).status_code == 200         # no-op
    assert api.patch(f"{V1}/contracts/{contract_id}",
                     json={"status": "SUPERSEDED"}).json()["data"]["status"] == "SUPERSEDED"
    assert api.patch(f"{V1}/contracts/{contract_id}",
                     json={"status": "DRAFT"}).json()["data"]["status"] == "DRAFT"
    assert api.patch(f"{V1}/contracts/{contract_id}",
                     json={"status": "EXECUTED"}).status_code == 422       # not a state

    events = db.execute(select(M.AuditEvent)
                        .where(M.AuditEvent.action == "contract.status_changed")
                        .order_by(M.AuditEvent.timestamp, M.AuditEvent.id)
                        ).scalars().all()
    # `timestamp` is now() — transaction start — so all three share it inside one
    # test transaction; the property is WHICH transitions were trailed, not order.
    assert sorted((e.before_state["status"], e.after_state["status"]) for e in events) == sorted([
        ("DRAFT", "ACTIVE"), ("ACTIVE", "SUPERSEDED"), ("SUPERSEDED", "DRAFT")])
    assert all(e.actor_id == owner.id and e.entity_id == uuid.UUID(contract_id)
               for e in events)


def test_a_failed_reread_changes_nothing_for_readers(api, db, owner, monkeypatch):
    """Adversarial: the current parser cannot read a file the old one could. The
    REPROCESS run is recorded FAILED (42.5 history), but the standing reading is
    still the document — same rows in the pane, same statuses on the version,
    same chunks in the index. A re-read may add a reading; it never takes one."""
    from legalmind.assist import store as assist_store
    from legalmind.ingestion import parsing

    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = uuid.UUID(_upload(api, contract_id))
    rows_before = {r["id"] for r in
                   api.get(f"{V1}/document-versions/{version_id}/evidence").json()["data"]}
    version = db.get(M.DocumentVersion, version_id)
    statuses_before = (version.processing_status, version.extraction_status)
    chunks_before = assist_store.count_chunks(db, version_id)

    def broken(*args, **kwargs):
        raise parsing.ParseError("simulated: the current parser cannot read this file")
    monkeypatch.setattr(parsing, "parse", broken)

    response = api.post(f"{V1}/document-versions/{version_id}/reprocess")
    assert response.status_code == 201, response.text
    assert response.json()["data"]["processing_run"]["status"] == "FAILED"
    assert response.json()["data"]["evidence_count"] == 0

    db.expire_all()
    version = db.get(M.DocumentVersion, version_id)
    assert (version.processing_status, version.extraction_status) == statuses_before
    assert {r["id"] for r in
            api.get(f"{V1}/document-versions/{version_id}/evidence").json()["data"]} == rows_before
    assert assist_store.count_chunks(db, version_id) == chunks_before


def test_a_second_concurrent_reread_is_refused_rather_than_starting_a_second_run(
        api, db, owner, engine):
    """The check-then-act race, closed (2026-09-06).

    `_reprocess_blockers` reading "nothing in flight" and the run being created
    are two steps; without a lock two callers — a double-click, a client retry,
    a script — could both pass the check and both create a REPROCESS run, and
    the assist index would then be built from whichever finished last rather
    than from the run every reader resolves to (`latest_completed_run_id`).

    The lock is held here on a REAL second connection to the SAME database,
    which is what a concurrent request — or the background OCR job, which takes
    the same `version_lock_key` — actually is. The endpoint must refuse with 409
    and write NOTHING. (Advisory locks are per-database, so this has to be the
    test engine: a lock taken on any other database would not conflict and the
    test would pass for the wrong reason.)
    """
    from sqlalchemy import text as sql_text

    from legalmind.worker.dispatch import version_lock_key

    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts",
                           json={"name": "ACME MSA"}).json()["data"]["id"]
    version_id = uuid.UUID(_upload(api, contract_id))

    def runs() -> int:
        return db.execute(
            select(func.count(M.DocumentProcessingRun.id))
            .where(M.DocumentProcessingRun.document_version_id == version_id)
        ).scalar_one()

    before = runs()
    key = version_lock_key(version_id)
    with engine.connect() as holder:
        assert holder.execute(sql_text("SELECT pg_try_advisory_lock(:k)"),
                              {"k": key}).scalar() is True

        refused = api.post(f"{V1}/document-versions/{version_id}/reprocess")
        assert refused.status_code == 409, refused.text
        assert "already being re-read" in refused.text
        assert runs() == before, "a refused re-read must create no processing run"

        # The holder releases; the very same request now succeeds — proving the
        # 409 was the lock and not some other refusal.
        holder.execute(sql_text("SELECT pg_advisory_unlock(:k)"), {"k": key})
        holder.commit()

    accepted = api.post(f"{V1}/document-versions/{version_id}/reprocess")
    assert accepted.status_code == 201, accepted.text
    assert runs() == before + 1


# =====================================================================
# AB-13 — the counterparty as an entity (2026-09-06)
# =====================================================================
def test_a_counterparty_groups_a_companys_documents_without_a_relationship_table(
        api, db, owner):
    """AB-13 r3 — the manager's "related documents" answer.

    Two contracts (an NDA and an MSA) linked to ONE company come back together
    from that company's own endpoint. The grouping is derived from the FK, not
    from any document-to-document relationship row — there is no such table and
    this proves none is needed."""
    sign_in(api, db, owner)
    created = api.post(f"{V1}/counterparties",
                       json={"name": "Placeholder Counterparty Ltd",
                             "industry": "Cloud hosting"})
    assert created.status_code == 201, created.text
    cp = created.json()["data"]
    assert cp["industry"] == "Cloud hosting"

    ids = []
    for name, doc_type in (("ACME NDA", "NDA"), ("ACME MSA", "MSA")):
        contract_id = api.post(f"{V1}/contracts",
                               json={"name": name, "contract_type": doc_type}
                               ).json()["data"]["id"]
        linked = api.patch(f"{V1}/contracts/{contract_id}",
                           json={"counterparty_id": cp["id"]})
        assert linked.status_code == 200, linked.text
        assert linked.json()["data"]["counterparty_id"] == cp["id"]
        ids.append(contract_id)

    profile = api.get(f"{V1}/counterparties/{cp['id']}").json()["data"]
    assert {c["id"] for c in profile["contracts"]} == set(ids)
    assert {c["contract_type"] for c in profile["contracts"]} == {"NDA", "MSA"}

    # The link is audited (r8) — it is the fact the whole view derives from.
    linked_events = db.execute(
        select(M.AuditEvent)
        .where(M.AuditEvent.action == "contract.counterparty_linked")).scalars().all()
    assert len(linked_events) == 2
    assert all(e.after_state["counterparty_id"] == cp["id"] for e in linked_events)


def test_an_unknown_industry_stays_absent_rather_than_becoming_a_blank(api, db, owner):
    """AB-13 r1 / rule 21 — a counterparty that is not fully known yet is the
    NORMAL case (the manager's seventh point). Nothing is invented, and the key
    is OMITTED rather than nulled so no UI can render "Industry: —" as though
    it had been checked."""
    sign_in(api, db, owner)
    bare = api.post(f"{V1}/counterparties", json={"name": "Unknown Co"}).json()["data"]
    assert bare["name"] == "Unknown Co"
    assert "industry" not in bare and "relationship_notes" not in bare

    # Whitespace is "nothing said", not an industry called " ".
    blank = api.post(f"{V1}/counterparties",
                     json={"name": "Blank Co", "industry": "   "}).json()["data"]
    assert "industry" not in blank

    # It can be filled in later, and cleared again.
    filled = api.patch(f"{V1}/counterparties/{bare['id']}",
                       json={"industry": "Logistics"}).json()["data"]
    assert filled["industry"] == "Logistics"
    cleared = api.patch(f"{V1}/counterparties/{bare['id']}",
                        json={"industry": None}).json()["data"]
    assert "industry" not in cleared
    # A name cannot be erased — a company with no name is no identity (r1).
    kept = api.patch(f"{V1}/counterparties/{bare['id']}", json={"name": None})
    assert kept.json()["data"]["name"] == "Unknown Co"


def test_a_counterparty_is_never_visible_outside_the_callers_contract_scope(
        api, db, owner):
    """AB-13 r6 — the disclosure boundary, and the reason there is no global list.

    "We have a deal with X" is exactly the class of fact `SEC-07`/`LEGAL-02` keep
    inside scope. Another user must not see this company at all: not in the list,
    and not by id — where the answer is the byte-identical 404 an absent row
    gets, never a 403 that would confirm it exists."""
    sign_in(api, db, owner)
    cp = api.post(f"{V1}/counterparties", json={"name": "Confidential Co"}).json()["data"]
    contract_id = api.post(f"{V1}/contracts", json={"name": "ACME MSA"}).json()["data"]["id"]
    api.patch(f"{V1}/contracts/{contract_id}", json={"counterparty_id": cp["id"]})
    assert [c["id"] for c in api.get(f"{V1}/counterparties").json()["data"]] == [cp["id"]]
    sign_out(api)

    stranger = make_user(db)
    grant_role(db, stranger, P.ROLE_USER)
    sign_in(api, db, stranger)
    assert api.get(f"{V1}/counterparties").json()["data"] == []
    assert api.get(f"{V1}/counterparties/{cp['id']}").status_code == 404
    missing = api.get(f"{V1}/counterparties/{uuid.uuid4()}")
    assert missing.status_code == 404
    # Indistinguishable: out-of-scope and nonexistent give the same code and the
    # same message (49.5 r1). `request_id` differs by design — it is the 49.9
    # correlation handle, and it carries no information about the object.
    out_of_scope = api.get(f"{V1}/counterparties/{cp['id']}").json()["error"]
    nonexistent = missing.json()["error"]
    assert out_of_scope["code"] == nonexistent["code"] == "NOT_FOUND"
    assert out_of_scope["message"] == nonexistent["message"]


def test_linking_refuses_an_unknown_counterparty_rather_than_dangling(api, db, owner):
    """A dangling link is worse than no link: the "everything for this company"
    view would silently lose the contract."""
    sign_in(api, db, owner)
    contract_id = api.post(f"{V1}/contracts", json={"name": "ACME MSA"}).json()["data"]["id"]
    refused = api.patch(f"{V1}/contracts/{contract_id}",
                        json={"counterparty_id": str(uuid.uuid4())})
    assert refused.status_code == 422
    assert db.get(M.Contract, uuid.UUID(contract_id)).counterparty_id is None

    # Unlinking is explicit and audited; left out entirely, the link is untouched.
    cp = api.post(f"{V1}/counterparties", json={"name": "Acme"}).json()["data"]
    api.patch(f"{V1}/contracts/{contract_id}", json={"counterparty_id": cp["id"]})
    api.patch(f"{V1}/contracts/{contract_id}", json={"name": "Renamed only"})
    assert str(db.get(M.Contract, uuid.UUID(contract_id)).counterparty_id) == cp["id"]
    unlinked = api.patch(f"{V1}/contracts/{contract_id}", json={"counterparty_id": None})
    assert unlinked.json()["data"]["counterparty_id"] is None


def test_a_counterparty_with_contracts_cannot_be_deleted_out_from_under_them(
        api, db, owner):
    """AB-13 r2 — ON DELETE RESTRICT. Nothing in the API deletes a counterparty;
    this pins the database's own refusal, so a future route cannot orphan a
    contract by accident."""
    import sqlalchemy

    sign_in(api, db, owner)
    cp = api.post(f"{V1}/counterparties", json={"name": "Acme"}).json()["data"]
    contract_id = api.post(f"{V1}/contracts", json={"name": "ACME MSA"}).json()["data"]["id"]
    api.patch(f"{V1}/contracts/{contract_id}", json={"counterparty_id": cp["id"]})
    db.flush()

    with pytest.raises(sqlalchemy.exc.IntegrityError):
        db.execute(sqlalchemy.text("DELETE FROM counterparties WHERE id = :i"),
                   {"i": uuid.UUID(cp["id"])})
        db.flush()
    db.rollback()


def test_no_endpoint_lists_every_counterparty(api, db, owner):
    """AB-13 r6 states the global list is FORBIDDEN, not merely unbuilt. Pinned
    structurally so a future convenience endpoint has to confront the record."""
    from legalmind.api.permission_map import ENDPOINT_PERMISSIONS

    counterparty_routes = {(m, p) for (m, p) in ENDPOINT_PERMISSIONS
                           if "/counterparties" in p}
    assert counterparty_routes == {
        ("GET", f"{V1}/counterparties"),
        ("POST", f"{V1}/counterparties"),
        ("GET", f"{V1}/counterparties/{{counterparty_id}}"),
        ("PATCH", f"{V1}/counterparties/{{counterparty_id}}"),
    }
    # And the two reads are contract-scoped permissions, not an admin one (r5).
    assert ENDPOINT_PERMISSIONS[("GET", f"{V1}/counterparties")] == P.CONTRACT_VIEW
    assert ENDPOINT_PERMISSIONS[("POST", f"{V1}/counterparties")] == P.CONTRACT_UPDATE


def test_a_lead_sees_a_companys_documents_across_the_department_and_the_shelf(
        api, db, owner):
    """Regression, found in the pre-commit review (2026-09-06).

    `_readable` grants sight of a company through a DEPARTMENT contract, but the
    detail listed only `own` — so a Department Lead opened the company the
    "everything for this company" view exists for and saw an EMPTY list. The
    same mismatch hid ARCHIVED contracts, which AB-12 r6 keeps rather than
    destroys. Both corrected: the list is now the caller's full read scope, in
    both archive states, and each row still carries `archived_at` so they are
    told apart."""
    from legalmind.db import models as M2

    dept = M2.Department(code=f"D{uuid.uuid4().hex[:6]}", name="Legal")
    db.add(dept); db.flush()
    owner.department_id = dept.id
    db.flush()

    sign_in(api, db, owner)
    cp = api.post(f"{V1}/counterparties", json={"name": "Shared Co"}).json()["data"]
    live = api.post(f"{V1}/contracts", json={"name": "Live deal"}).json()["data"]["id"]
    shelved = api.post(f"{V1}/contracts", json={"name": "Old deal"}).json()["data"]["id"]
    for cid in (live, shelved):
        api.patch(f"{V1}/contracts/{cid}", json={"counterparty_id": cp["id"]})
    api.post(f"{V1}/contracts/{shelved}/archive")
    sign_out(api)

    # A colleague in the SAME department, holding the Lead's read scope.
    lead = make_user(db)
    grant(db, lead, bespoke_role(db, f"LEAD{uuid.uuid4().hex[:4]}",
                                 [P.CONTRACT_VIEW, P.DEPARTMENT_VIEW]))
    lead.department_id = dept.id
    db.flush()
    sign_in(api, db, lead)

    profile = api.get(f"{V1}/counterparties/{cp['id']}")
    assert profile.status_code == 200, profile.text
    listed = profile.json()["data"]["contracts"]
    assert {c["id"] for c in listed} == {live, shelved}, (
        "the Lead must see the department's deals for this company, archived included")
    assert [c["archived_at"] is not None for c in listed].count(True) == 1

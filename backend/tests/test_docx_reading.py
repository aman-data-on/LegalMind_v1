"""DOCX reading in document order, every story read once — 2026-10-15.

Findings F2a–F2e, PB, F4, N-HIDDEN and N-SILENT (document-reading verification) and
BR-4 (headers and footers are page furniture, never a clause). Before this change a
DOCX table came after all body text with no page, its cells joined by tabs that
normalization collapsed; merged cells repeated, nested tables vanished; text boxes,
footnotes, endnotes, custom XML, rows in content controls, headers and footers were
never read; hidden text was; and a document with comments said nothing about them.

Every fixture is built in memory with python-docx — no real document enters the
repository (locked 54.6).
"""

from __future__ import annotations

import io
import re
import zipfile
from types import SimpleNamespace

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement

from legalmind.domain import enums as E
from legalmind.ingestion.parsing import parse_docx

_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml."


def _save(d) -> bytes:
    buffer = io.BytesIO()
    d.save(buffer)
    return buffer.getvalue()


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _rows(result) -> list:
    """Segments as evidence-like rows, for the pure consumers below."""
    return [SimpleNamespace(id=f"e{i}", content=s.content, section_number=s.section_number,
                            section_title=s.section_title, page_number=s.page_number,
                            start_offset=s.start_offset, end_offset=s.end_offset,
                            source_type=s.source_type, evidence_metadata=s.metadata)
            for i, s in enumerate(result.segments)]


def test_a_table_is_read_where_it_stands_one_row_per_line_each_cell_once() -> None:
    """F2a–F2e: the table between 7.1 and 7.2 comes between them, cells survive
    normalization as ` | `, a gridSpan cell reads once, a vMerge cell's value is on
    every row it spans, a two-paragraph cell stays one row, a nested table is kept."""
    d = docx.Document()
    d.add_paragraph("7.1 Credits are as set out in the table below.")
    t = d.add_table(rows=4, cols=3)
    for col, text in enumerate(("Availability", "Credit", "Cap")):
        t.cell(0, col).text = text
    for row, (avail, credit) in enumerate((("< 99.9%", "5%"), ("< 99.0%", "10%"),
                                           ("< 95.0%", "20%")), start=1):
        t.cell(row, 0).text, t.cell(row, 1).text = avail, credit
    t.cell(1, 2).merge(t.cell(3, 2)).text = "Max 20% of monthly fee"
    t.cell(2, 1).add_paragraph("subject to clause 9")
    span = t.add_row()
    span.cells[0].merge(span.cells[2]).text = "Credits exclude Scheduled Maintenance"
    nested = t.cell(1, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).text, nested.cell(0, 1).text = "NESTED-A", "NESTED-B"
    d.add_paragraph("7.2 Credits are the sole remedy.")

    result = parse_docx(_save(d))
    kinds = [s.source_type for s in result.segments]
    assert kinds == [E.EvidenceSourceType.NATIVE_TEXT, E.EvidenceSourceType.TABLE,
                     E.EvidenceSourceType.NATIVE_TEXT]
    starts = [s.start_offset for s in result.segments]
    assert starts == sorted(starts) and len(set(starts)) == 3
    table = result.segments[1]
    assert table.content.splitlines() == [
        "| Availability | Credit | Cap |",
        "| < 99.9% [NESTED-A | NESTED-B] | 5% | Max 20% of monthly fee |",
        "| < 99.0% | 10% subject to clause 9 | Max 20% of monthly fee |",
        "| < 95.0% | 20% | Max 20% of monthly fee |",
        "| Credits exclude Scheduled Maintenance |",
    ]
    assert table.metadata == {"table_index": 1, "rows": 5}
    assert table.section_number is None            # numbering is never inherited


def test_a_break_inside_a_table_cell_moves_every_later_page() -> None:
    """PB: the break in row 3 put '3. Liability' on page 1 and pages_total at 1 short;
    the table takes the page of its first cell paragraph."""
    d = docx.Document()
    d.add_paragraph("1. Definitions on page one.")
    t = d.add_table(rows=3, cols=2)
    for i in range(3):
        t.cell(i, 0).text, t.cell(i, 1).text = f"Tier {i}", f"{i * 5}%"
    t.cell(2, 0).paragraphs[0].runs[0].add_break(WD_BREAK.PAGE)
    d.add_paragraph("3. Liability on page two.")

    result = parse_docx(_save(d))
    assert [(s.source_type, s.page_number) for s in result.segments] == [
        (E.EvidenceSourceType.NATIVE_TEXT, 1), (E.EvidenceSourceType.TABLE, 1),
        (E.EvidenceSourceType.NATIVE_TEXT, 2)]
    assert result.pages_total == 2


def test_a_page_break_inside_a_paragraph_does_not_fuse_its_words() -> None:
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Request within sixty (60) calendar days.")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run("Issued within forty-five (45) days.")
    [segment] = parse_docx(_save(d)).segments
    assert "days.Issued" not in segment.content
    assert "days. Issued" in " ".join(segment.content.split())


def test_hidden_text_is_not_the_document_s_text() -> None:
    """N-HIDDEN: a w:vanish run is neither displayed nor printed."""
    d = docx.Document()
    p = d.add_paragraph("3. Term. Visible text.")
    p.add_run(" HIDDEN-TEXT liability is unlimited.")._r.get_or_add_rPr().append(
        OxmlElement("w:vanish"))
    [segment] = parse_docx(_save(d)).segments
    assert segment.content == "3. Term. Visible text."


def _every_story() -> bytes:
    """A DOCX carrying each story the reader used to miss — the verification's f4
    builder: header, footer, footnote, endnote, a comment, a VML and a DrawingML text
    box (Word writes the latter twice, once in mc:Fallback), body custom XML and a
    table row inside a content control."""
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "HDR-TEXT Confidential - Draft v3"
    d.sections[0].footer.paragraphs[0].text = "FTR-TEXT Page footer Initials"
    d.add_paragraph("1. Fees. The Customer shall pay the Fees.")
    d.add_paragraph("2. Liability. See box.")
    d.add_paragraph("3. Term.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "ROW1-A", "ROW1-B"
    d.add_paragraph("4. Law.")
    files = {}
    with zipfile.ZipFile(io.BytesIO(_save(d))) as z:
        files = {n: z.read(n) for n in z.namelist()}
    doc = files["word/document.xml"].decode()
    doc = doc.replace(
        "1. Fees. The Customer shall pay the Fees.</w:t></w:r>",
        "1. Fees. The Customer shall pay the Fees.</w:t></w:r>"
        '<w:r><w:footnoteReference w:id="1"/></w:r><w:r><w:endnoteReference w:id="1"/></w:r>'
        '<w:commentRangeStart w:id="0"/>')
    vml = ('<w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>'
           f'<w:txbxContent><w:p>{_run("VML-TEXTBOX Liability capped at 3 months fees")}'
           '</w:p></w:txbxContent></v:textbox></v:shape></w:pict></w:r>')
    box = f'<w:txbxContent><w:p>{_run("DML-TEXTBOX Exclusions apply")}</w:p></w:txbxContent>'
    dml = ('<w:r><mc:AlternateContent '
           'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
           '<mc:Choice Requires="wps"><w:drawing><wp:anchor xmlns:wp="http://schemas.'
           'openxmlformats.org/drawingml/2006/wordprocessingDrawing"><a:graphic xmlns:a='
           '"http://schemas.openxmlformats.org/drawingml/2006/main"><a:graphicData uri="x">'
           '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/'
           f'wordprocessingShape"><wps:txbx>{box}</wps:txbx></wps:wsp></a:graphicData>'
           '</a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape '
           f'xmlns:v="urn:schemas-microsoft-com:vml"><v:textbox>{box}</v:textbox></v:shape>'
           '</w:pict></mc:Fallback></mc:AlternateContent></w:r>')
    doc = doc.replace("2. Liability. See box.</w:t></w:r>",
                      "2. Liability. See box.</w:t></w:r>" + vml + dml)
    doc = doc.replace("3. Term.</w:t></w:r></w:p>",
                      '3. Term.</w:t></w:r></w:p><w:customXml w:element="x"><w:p>'
                      + _run("CUSTOMXML-PARA Renewal is automatic.") + "</w:p></w:customXml>")
    doc = re.sub(r"(</w:tr>)(</w:tbl>)",
                 r"\1<w:sdt><w:sdtContent><w:tr><w:tc><w:p>" + _run("SDT-ROW-A")
                 + r"</w:p></w:tc><w:tc><w:p>" + _run("SDT-ROW-B")
                 + r"</w:p></w:tc></w:tr></w:sdtContent></w:sdt>\2", doc, count=1)
    doc = doc.replace("4. Law.</w:t></w:r>", '4. Law.</w:t></w:r><w:commentRangeEnd w:id="0"/>'
                      '<w:r><w:commentReference w:id="0"/></w:r>')
    files["word/document.xml"] = doc.encode()
    separator = '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
    files["word/footnotes.xml"] = (
        f'<w:footnotes {_W}>{separator}<w:footnote w:id="1"><w:p>'
        f'{_run("FOOTNOTE-TEXT Fees exclude GST at 18%")}</w:p></w:footnote>'
        f'<w:footnote w:id="2"><w:p>{_run("ORPHAN-NOTE never referenced")}</w:p></w:footnote>'
        "</w:footnotes>").encode()
    files["word/endnotes.xml"] = (
        f'<w:endnotes {_W}><w:endnote w:id="1"><w:p>{_run("ENDNOTE-TEXT Rates revised")}'
        "</w:p></w:endnote></w:endnotes>").encode()
    files["word/comments.xml"] = (
        f'<w:comments {_W}><w:comment w:id="0" w:author="x"><w:p>{_run("COMMENT-TEXT note")}'
        "</w:p></w:comment></w:comments>").encode()
    rels = files["word/_rels/document.xml.rels"].decode().replace("</Relationships>", "".join(
        f'<Relationship Id="rIdX{n}" Type="{_RELS}{n}" Target="{n}.xml"/>'
        for n in ("footnotes", "endnotes", "comments")) + "</Relationships>")
    files["word/_rels/document.xml.rels"] = rels.encode()
    files["[Content_Types].xml"] = files["[Content_Types].xml"].decode().replace(
        "</Types>", "".join(f'<Override PartName="/word/{n}.xml" ContentType="{_CT}{n}+xml"/>'
                            for n in ("footnotes", "endnotes", "comments")) + "</Types>").encode()
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, blob in files.items():
            z.writestr(name, blob)
    return out.getvalue()


def test_every_story_is_read_once_in_place_and_comments_are_disclosed_not_read() -> None:
    """F4 + N-SILENT."""
    result = parse_docx(_every_story())
    contents = [s.content for s in result.segments]
    text = "\n".join(contents)
    for marker in ("FOOTNOTE-TEXT", "ENDNOTE-TEXT", "VML-TEXTBOX", "DML-TEXTBOX",
                   "CUSTOMXML-PARA", "SDT-ROW-A", "HDR-TEXT", "FTR-TEXT"):
        assert text.count(marker) == 1, marker
    assert "COMMENT-TEXT" not in text
    # a note nothing visible refers to is not text the document shows (review 2026-10-16)
    assert "ORPHAN-NOTE" not in text
    # owner, 2026-10-27: a comment is kept beside its clause, never as the clause's text
    assert "1 of 1 comment(s) kept beside their clause (reviewer notes, never contract " \
           "text)" in result.diagnostics
    noted = [s for s in result.segments if s.metadata.get("comments")]
    assert len(noted) == 1 and noted[0].metadata["comments"] == [
        {"author": "x", "text": "COMMENT-TEXT note"}]
    assert not any("comment_ids" in s.metadata for s in result.segments)

    first = [c.split()[0] for c in contents]
    assert first == ["1.", "FOOTNOTE-TEXT", "ENDNOTE-TEXT", "2.", "VML-TEXTBOX",
                     "DML-TEXTBOX", "3.", "CUSTOMXML-PARA", "|", "4.",
                     "HDR-TEXT", "FTR-TEXT"]
    starts = [s.start_offset for s in result.segments]
    assert starts == sorted(starts)
    by_marker = {c.split()[0]: s for c, s in zip(contents, result.segments, strict=True)}
    assert by_marker["FOOTNOTE-TEXT"].metadata == {"part": "footnote", "note_id": "1"}
    assert by_marker["ENDNOTE-TEXT"].metadata == {"part": "endnote", "note_id": "1"}
    for marker, part in (("HDR-TEXT", "header"), ("FTR-TEXT", "footer")):
        segment = by_marker[marker]
        assert segment.source_type is E.EvidenceSourceType.OTHER
        assert segment.metadata == {"part": part} and segment.page_number is None
    assert "| ROW1-A | ROW1-B |\n| SDT-ROW-A | SDT-ROW-B |" in contents
    assert result.status is E.ExtractionStatus.COMPLETE


_HEADER = ("Leapswitch Networks Private Limited - Master Services Agreement - "
           "Confidential - Draft for discussion")


def _three_sections() -> bytes:
    """Three sections: the second's header linked to the first (python-docx returns
    the linked header for it), the third's its own copy of the same text, plus a
    first-page variant with that text too — three copies unless deduplicated."""
    d = docx.Document()
    assert len(_HEADER) >= 100
    d.sections[0].header.paragraphs[0].text = _HEADER
    d.sections[0].different_first_page_header_footer = True
    d.sections[0].first_page_header.paragraphs[0].text = _HEADER
    d.sections[0].footer.paragraphs[0].text = "EXHIBIT A - PRICING"
    for n in (1, 2, 3):
        if n > 1:
            section = d.add_section()
            if n == 3:
                section.header.is_linked_to_previous = False
                section.header.paragraphs[0].text = _HEADER
        d.add_paragraph(f"{n}. Clause {n}. The parties agree to the terms set out in "
                        f"this clause number {n} of the agreement between them.")
    return _save(d)


def test_a_linked_header_is_read_once_and_does_not_stop_indexing() -> None:
    """BR-4: a header repeated per section tripped REPEATED_TEXT, and the document
    then silently stopped being indexed."""
    from legalmind.assist.ingestion.chunking import chunk_evidence, integrity_failures

    result = parse_docx(_three_sections())
    furniture = [s for s in result.segments if s.metadata.get("part")]
    assert [(s.metadata["part"], s.content) for s in furniture] == [
        ("header", _HEADER), ("footer", "EXHIBIT A - PRICING")]
    rows = _rows(result)
    assert integrity_failures(rows, chunk_evidence(rows)) == []


def test_headers_and_footers_never_become_a_clause() -> None:
    """BR-4: an 'EXHIBIT A' footer row opened a bogus Diff-Talk part and stopped
    clause splitting; the shared filter keeps every clause reader's input as before."""
    from legalmind.analysis.document_comparison import version_clauses
    from legalmind.mapping.service import body_rows

    rows = _rows(parse_docx(_three_sections()))
    assert any(r.evidence_metadata.get("part") == "footer" for r in rows)
    assert [c["ref"] for c in version_clauses(body_rows(rows))] == ["1", "2", "3"]
    note = SimpleNamespace(evidence_metadata={"part": "footnote"})
    assert body_rows([note]) == [note]                # a footnote IS body content


def test_load_clauses_leaves_headers_and_footers_out(db, tmp_path) -> None:
    """BR-4 (c): every evidence row was a mapping candidate, so a header could be
    confirmed as a presence clause. Footnotes stay — they are body content."""
    from legalmind.db import models as M
    from legalmind.ingestion.service import ingest_document
    from legalmind.ingestion.storage import LocalFilesystemStorage
    from legalmind.ingestion.validation import DOCX_MIME
    from legalmind.mapping.service import load_clauses
    from tests.conftest import make_user

    owner = make_user(db)
    contract = M.Contract(owner_id=owner.id, name="ACME MSA",
                          status=E.ContractStatus.ACTIVE)
    db.add(contract); db.flush()
    result = ingest_document(
        db, LocalFilesystemStorage(tmp_path / "objects"), contract_id=contract.id,
        uploaded_by=owner.id, data=_every_story(), filename="msa.docx",
        declared_mime=DOCX_MIME)
    stored = db.query(M.DocumentEvidence).filter_by(
        processing_run_id=result.processing_run.id).all()
    assert {e.evidence_metadata.get("part") for e in stored} >= {"header", "footer"}
    text = "\n".join(c.content for c in load_clauses(db, result.document_version.id))
    assert "HDR-TEXT" not in text and "FTR-TEXT" not in text
    assert "FOOTNOTE-TEXT" in text and "SDT-ROW-A" in text


def _rewritten(data: bytes, **parts: bytes | str) -> bytes:
    """``data`` with zip members replaced (``word_document`` -> word/document.xml)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        files = {n: z.read(n) for n in z.namelist()}
    for name, blob in parts.items():
        files[name.replace("_", "/", 1) + ".xml"] = blob.encode() if isinstance(
            blob, str) else blob
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, blob in files.items():
            z.writestr(name, blob)
    return out.getvalue()


def test_a_note_or_text_box_the_reader_cannot_see_is_never_read() -> None:
    """Review 2026-10-16: a footnote whose reference sits in hidden text or a tracked
    deletion, and a text box anchored in hidden text, reached the evaluator's input —
    text the accepted, visible document does not contain."""
    base = _every_story()
    with zipfile.ZipFile(io.BytesIO(base)) as z:
        doc = z.read("word/document.xml").decode()
    vanish = '<w:rPr><w:vanish/></w:rPr>'
    doc = doc.replace('<w:r><w:footnoteReference w:id="1"/></w:r>',
                      f'<w:r>{vanish}<w:footnoteReference w:id="1"/></w:r>')
    doc = doc.replace('<w:r><w:endnoteReference w:id="1"/></w:r>',
                      '<w:del w:id="9" w:author="x"><w:r><w:endnoteReference w:id="1"/>'
                      '</w:r></w:del>')
    doc = doc.replace("<w:r><w:pict>", f"<w:r>{vanish}<w:pict>", 1)
    text = "\n".join(s.content for s in parse_docx(_rewritten(
        base, word_document=doc)).segments)
    for marker in ("FOOTNOTE-TEXT", "ENDNOTE-TEXT", "VML-TEXTBOX"):
        assert marker not in text, marker
    assert "DML-TEXTBOX" in text and "1. Fees." in text      # the visible ones still are


def test_a_row_split_across_a_page_is_one_boundary_not_one_per_cell() -> None:
    """Review 2026-10-16: Word records the boundary of a row split across pages in EVERY
    cell; counted per cell, a two-cell row moved every later paragraph two pages on."""
    d = docx.Document()
    d.add_paragraph("1. Definitions on page one.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Tier credit", "Fifteen percent"
    for col in (0, 1):
        t.cell(0, col).paragraphs[0].runs[0]._r.append(
            OxmlElement("w:lastRenderedPageBreak"))
    d.add_paragraph("3. Liability on page two.")
    result = parse_docx(_save(d))
    assert result.segments[-1].page_number == 2 and result.pages_total == 2


def test_a_damaged_footnotes_part_is_a_diagnostic_not_a_failed_document() -> None:
    """Review 2026-10-16: a broken auxiliary part raised past the parser and failed the
    whole upload; the body is still the document's text."""
    data = _rewritten(_every_story(), word_footnotes=b"<w:footnotes <<not xml")
    result = parse_docx(data)
    # read, but not whole: the damaged part may hold the clause asked about (rule 15)
    assert result.status is E.ExtractionStatus.PARTIAL
    assert any(d.startswith("footnotes and endnotes could not be read")
               for d in result.diagnostics)
    assert "1. Fees." in result.segments[0].content


def test_imported_content_it_cannot_read_makes_the_reading_partial() -> None:
    """Review 2026-10-16: an altChunk (another document imported whole) was counted in
    diagnostics while the reading stayed COMPLETE."""
    d = docx.Document()
    d.add_paragraph("1. Fees. The Customer shall pay the Fees.")
    d.element.body.insert(1, OxmlElement("w:altChunk"))
    result = parse_docx(_save(d))
    assert result.status is E.ExtractionStatus.PARTIAL
    assert "1 embedded object(s) not read" in result.diagnostics

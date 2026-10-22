"""Legal regression set (owner, 2026-10-21): realistic multi-clause contracts run through
the real pipeline — ingest → map → extract → evaluate → angles — asserting BOTH the
evidence (which clause, which page) and the result (classification, action, the business
line, the observations). Synthetic shapes only (locked 54.6, rule 21): no real
counterparty; every expectation names a ratified standard by code and never copies its
value. A known gap is `xfail(strict=True)` with its root cause, so the fix flips it.
"""
from __future__ import annotations

import io
import os
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK
from sqlalchemy import select

from legalmind.analysis.service import run_analysis
from legalmind.api.angles import review_angles
from legalmind.db import models as M
from legalmind.domain import enums as E
from legalmind.domain.constitution_boundaries import constitution_prohibition_for
from legalmind.ingestion.service import ingest_document
from tests.test_analysis import DOCX_MIME, build, storage  # noqa: F401 — fixtures
from tests.test_hard_test_fixes import _ratified

PAGE = object()          # a page break the file itself states (w:br type="page")


def _docx(paragraphs: list) -> bytes:
    document = docx.Document()
    for item in paragraphs:
        if item is PAGE:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            document.add_paragraph(item)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _review(build, db, codes: list[str], paragraphs: list, kind: str = "MSA"):
    for code in codes:
        _ratified(build, code)
    contract = M.Contract(owner_id=build.owner.id, name="Regression contract",
                          contract_type=kind, status=E.ContractStatus.ACTIVE)
    db.add(contract); db.flush()
    result = ingest_document(db, build.storage, contract_id=contract.id,
                             uploaded_by=build.owner.id, data=_docx(paragraphs),
                             filename="contract.docx", declared_mime=DOCX_MIME)
    review = M.Review(contract_id=contract.id, document_version_id=result.document_version.id,
                      configuration_snapshot_id=build.snapshot.id,
                      status=E.ReviewStatus.DRAFT, created_by=build.owner.id)
    db.add(review); db.flush()
    run = run_analysis(db, review)
    return review, run


def _findings(db, review) -> dict[str, dict]:
    """code → classification, actual value, and the (clause, page) of every cited passage."""
    out: dict[str, dict] = {}
    rows = db.execute(
        select(M.Requirement.code, M.Finding)
        .join(M.RequirementVersion, M.RequirementVersion.id == M.Finding.requirement_version_id)
        .join(M.Requirement, M.Requirement.id == M.RequirementVersion.requirement_id)
        .where(M.Finding.review_id == review.id)).all()
    for code, finding in rows:
        cited = db.execute(
            select(M.DocumentEvidence.section_number, M.DocumentEvidence.page_number)
            .join(M.FindingEvidence, M.FindingEvidence.evidence_id == M.DocumentEvidence.id)
            .where(M.FindingEvidence.finding_id == finding.id)).all()
        actual = db.execute(select(M.Evaluation.actual_value)
                            .where(M.Evaluation.finding_id == finding.id)).scalars().first()
        out[code] = {"classification": finding.classification.value, "actual": actual or {},
                     "cited": sorted({(s, p) for s, p in cited}, key=str)}
    return out


def _clauses(found: dict) -> set:
    return {section for section, _ in found["cited"]}


def _angles(db, review) -> dict:
    return review_angles(db, review, legal_position=True)


# ------------------------------------------------------------------ exit cluster
FIXED_TERM_EXIT = [
    "1. Term",
    "1.1 Contract Term: twelve (12) months from the Effective Date.",
    "2. Fees",
    "2.1 The Customer shall pay Rs. 1,00,000 per month, exclusive of GST.",
    PAGE,
    "13. Service Termination",
    "13.1 Either party may terminate for convenience by providing fifteen (15) days' "
    "written notice to the other party.",
    "13.2 There is no lock-in period and no early termination charges shall apply.",
    "13.3 Any prepaid fees for the unutilised period shall be refunded to the Customer on a "
    "pro-rata basis within fifteen (15) days of termination.",
]
EXIT_CODES = ["CONVENIENCE-NOTICE-MSA-001", "EARLY-TERM-RESTRICTION-MSA-001"]


def test_exit_notice_is_measured_on_its_own_clause_and_page(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    found = _findings(db, review)
    notice = found["CONVENIENCE-NOTICE-MSA-001"]
    assert notice["classification"] == "DEVIATION"
    assert ("13.1", 2) in notice["cited"]


def test_exit_waiver_of_the_lock_in_is_a_deviation_citing_the_waiver(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    early = _findings(db, review)["EARLY-TERM-RESTRICTION-MSA-001"]
    assert early["classification"] == "DEVIATION"
    assert "13.2" in _clauses(early)


def test_exit_business_point_carries_the_refund_and_the_unsecured_value(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    angles = _angles(db, review)
    exit_point = next(p for p in angles["business"]
                      if p["requirement_code"] == "EARLY-TERM-RESTRICTION-MSA-001")
    assert "pro-rata" in exit_point["body"] and "15 days' notice" in exit_point["body"]
    assert exit_point["money"] and "12,00,000" in exit_point["money"]


def test_exit_notice_finding_is_read_with_the_charges_and_refund_clauses(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    item = next(i for i in _angles(db, review)["legal"]["items"]
                if i["requirement_code"] == "CONVENIENCE-NOTICE-MSA-001")
    related = {r["clause"]: r for r in item["related"]}
    assert set(related) == {"13.2", "13.3"}
    assert related["13.3"]["page"] == 2 and related["13.3"]["relation"] == (
        "prepaid fees on leaving")
    assert "13.3 Any prepaid fees for the unutilised period" in related["13.3"]["text"]
    assert "no early termination charges" in related["13.2"]["text"]


def test_a_termination_notice_stated_two_ways_is_observed_naming_both(build, db):
    paragraphs = [*FIXED_TERM_EXIT[:7], "20. Miscellaneous",
                  "20.1 Either party may terminate for convenience by providing thirty (30) "
                  "days' written notice to the other party."]
    review, _ = _review(build, db, EXIT_CODES, paragraphs, kind="OTHER")
    notice = _observed(db, review, "NOTICE_TWO_WAYS")
    assert len(notice) == 1 and notice[0]["clauses"] == ["13.1", "20.1"]


# ------------------------------------------------------------------ liability cluster
CAP = "LIABILITY-MSA-001"


def test_a_six_month_cap_deviates_and_cites_its_clause_and_page(build, db):
    review, _ = _review(build, db, [CAP], [
        "1. Services", "1.1 Leapswitch shall provide the Services.", PAGE,
        "17. Limitation of Liability",
        "17.2 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the six (6) months preceding the claim."])
    cap = _findings(db, review)[CAP]
    assert cap["classification"] == "DEVIATION"
    assert cap["cited"] == [("17.2", 2)]


def test_two_caps_for_one_scope_conflict_and_both_clauses_are_cited(build, db):
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability",
        "17.2 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the six (6) months preceding the claim.",
        "17.7 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the twelve (12) months preceding the claim."])
    cap = _findings(db, review)[CAP]
    assert cap["classification"] == "CONFLICT"
    assert {"17.2", "17.7"} <= _clauses(cap)


ONE_SIDED = ("17.2 Leapswitch's total aggregate liability shall not exceed the total fees "
             "paid by the Customer in the twelve (12) months preceding the claim.")


def _cap_case(build, db, sentence, codes=(CAP,), kind="MSA"):
    review, _ = _review(build, db, list(codes), ["17. Limitation of Liability", sentence],
                        kind=kind)
    return review, _findings(db, review)


def test_a_cap_binding_one_party_is_never_acceptable(build, db):
    """G-D (2026-10-21): Constitution §9 — "a cap applying to only one party" is an
    Unacceptable Position; the figure matching says nothing about whom it binds."""
    review, found = _cap_case(build, db, ONE_SIDED)
    cap = found[CAP]
    assert cap["cited"] == [("17.2", None)]
    assert cap["classification"] == "DEVIATION" and cap["actual"]["cap_party"] == "ours"
    assert constitution_prohibition_for(CAP, cap["actual"]) == {
        "section": "9", "quote": "[…] a cap applying to only one party."}
    item = next(i for i in _angles(db, review)["legal"]["items"]
                if i["requirement_code"] == CAP)
    assert item["action"].startswith("Must change before signing")


def test_a_cap_on_the_customer_only_is_one_sided_too(build, db):
    _, found = _cap_case(build, db, ONE_SIDED.replace("Leapswitch's", "The Customer's"))
    assert (found[CAP]["classification"], found[CAP]["actual"]["cap_party"]) == (
        "DEVIATION", "theirs")


def test_a_mutual_twelve_month_cap_still_matches(build, db):
    _, found = _cap_case(build, db, ONE_SIDED.replace("Leapswitch's", "Each party's"))
    assert found[CAP]["classification"] == "MATCH"
    assert found[CAP]["actual"]["cap_party"] == "both"


def test_a_cap_whose_subject_names_no_party_is_never_guessed(build, db):
    _, found = _cap_case(build, db, ONE_SIDED.replace("Leapswitch's total aggregate", "Total"))
    assert found[CAP]["classification"] == "MATCH" and "cap_party" not in found[CAP]["actual"]


def test_our_published_tos_cap_is_not_held_to_a_mutual_cap(build, db):
    _, found = _cap_case(build, db, ONE_SIDED, codes=("LIABILITY-TOS-001",), kind="TOS")
    assert found["LIABILITY-TOS-001"]["classification"] == "MATCH"


def test_a_deviating_cap_is_read_with_the_data_loss_exclusion(build, db):
    review, _ = _review(build, db, [CAP, "DATA-LOSS-EXCLUSION-MSA-001"], [
        "17. Limitation of Liability",
        "17.2 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the six (6) months preceding the claim.",
        "17.4 Leapswitch shall not be liable for any loss, deletion, corruption or "
        "unavailability of customer data, or the costs of restoring it."])
    item = next(i for i in _angles(db, review)["legal"]["items"]
                if i["requirement_code"] == CAP)
    related = {r["clause"]: r["relation"] for r in item["related"]}
    assert related.get("17.4") == "the data-loss exclusion"


def test_a_data_loss_exclusion_in_other_words_is_found(build, db):
    """G-E (2026-10-21): "shall not BE liable" never matched "not liable"."""
    review, _ = _review(build, db, ["DATA-LOSS-EXCLUSION-MSA-001"], [
        "17. Limitation of Liability",
        "17.4 Leapswitch shall not be liable for any loss, deletion or corruption of customer "
        "data, or the costs of restoring it."])
    assert _findings(db, review)["DATA-LOSS-EXCLUSION-MSA-001"]["classification"] == "MATCH"


def test_a_carve_out_that_wipes_out_the_cap_is_never_accepted(build, db):
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability",
        "17.2 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the twelve (12) months preceding the claim.",
        "17.3 The limitation in clause 17.2 shall not apply to any breach of this Agreement."])
    cap = _findings(db, review)[CAP]
    assert cap["classification"] != "MATCH"
    assert {"17.2", "17.3"} <= _clauses(cap)


def test_service_credits_in_an_sla_are_never_a_liability_cap(build, db):
    review, _ = _review(build, db, [CAP, "LIABILITY-TOS-001"], [
        "4. Service Credits",
        "4.1 Service credits shall not exceed twenty percent (20%) of the monthly fee."],
        kind="SLA")
    assert not {CAP, "LIABILITY-TOS-001"} & set(_findings(db, review))


# ------------------------------------------------------------------ recall on common wording
# Found by this set (2026-10-21, gaps G-A/G-B/G-C): drafting a lawyer reads at once that the
# lexical rules missed, so the review failed closed to "a person must read it".
def test_terminate_this_agreement_for_convenience_is_recognised(build, db):
    review, _ = _review(build, db, ["CONVENIENCE-NOTICE-MSA-001"], [
        "13. Termination", "13.1 Either party may terminate this agreement for convenience by "
        "providing fifteen (15) days' written notice."])
    assert _findings(db, review)["CONVENIENCE-NOTICE-MSA-001"]["classification"] == "DEVIATION"


def test_a_fixed_term_of_twelve_months_is_a_commitment(build, db):
    _, run = _review(build, db, ["EARLY-TERM-RESTRICTION-MSA-001"], [
        "1. Term", "1.1 This Agreement is for a fixed term of twelve (12) months.",
        "13. Termination", "13.2 There is no lock-in and no early termination charges shall "
        "apply."], kind="OTHER")
    coverage = {c["code"]: c["outcome"] for c in run.applicability}
    assert coverage["EARLY-TERM-RESTRICTION-MSA-001"] == "APPLIED"


def test_fees_paid_in_the_months_before_the_claim_is_a_fees_paid_basis(build, db):
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability", "17.2 The total aggregate liability of each party shall "
        "not exceed the fees paid in the six (6) months preceding the claim."])
    assert _findings(db, review)[CAP]["classification"] == "DEVIATION"


def test_two_bases_in_one_cap_sentence_still_fail_closed(build, db):
    """FEES_PAID and the affected-services basis stay distinct (owner, 2026-08-18)."""
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability", "17.2 The total aggregate liability of each party shall "
        "not exceed the fees paid in the six (6) months preceding the claim for the specific "
        "services giving rise to the claim."])
    assert _findings(db, review)[CAP]["classification"] == "UNABLE_TO_EVALUATE"


def test_a_fixed_term_with_no_length_is_not_a_commitment(build, db):
    _, run = _review(build, db, ["EARLY-TERM-RESTRICTION-MSA-001"], [
        "1. Term", "1.1 This Agreement is for a fixed term as set out in the Order Form."],
        kind="OTHER")
    coverage = {c["code"]: c["outcome"] for c in run.applicability}
    assert coverage["EARLY-TERM-RESTRICTION-MSA-001"] == "NOT_APPLICABLE"


# ------------------------------------------------------------------ renewal
def _observed(db, review, kind):
    observed = _angles(db, review)["observations"]
    return [o for o in observed["structural"] + observed["additional"] if o["kind"] == kind]


def test_renewal_stated_three_ways_is_one_observation_naming_all_three(build, db):
    review, _ = _review(build, db, ["AUTORENEW-MSA-001"], [
        "1. Definitions",
        "1.7 Renewal Term means any renewal period subject to the mutual written agreement of "
        "the Parties.",
        "5. Term",
        "5.2 This Agreement shall automatically renew for successive periods of one (1) year.",
        "5.3 Any renewal of this Agreement shall be subject to the mutual written agreement of "
        "the Parties."])
    renewal = _observed(db, review, "RENEWAL_TWO_WAYS")
    assert len(renewal) == 1 and set(renewal[0]["clauses"]) == {"1.7", "5.2", "5.3"}


def test_a_non_renewal_notice_stated_two_ways_is_observed(build, db):
    review, _ = _review(build, db, ["AUTORENEW-MSA-001"], [
        "5. Term",
        "5.2 This Agreement shall automatically renew unless either party gives notice of "
        "non-renewal at least thirty (30) days before the end of the current term.",
        "5.3 Either party may give written notice of its intention not to renew at least sixty "
        "(60) days prior to the end of the term."])
    notice = _observed(db, review, "NOTICE_TWO_WAYS")
    assert len(notice) == 1 and notice[0]["clauses"] == ["5.2", "5.3"]
    assert "30 days" in notice[0]["text"] and "60 days" in notice[0]["text"]


def test_a_customers_notice_and_the_providers_are_never_one_stated_twice(build, db):
    """MSA2 §25.1/§25.2: two rights, two periods — no contradiction."""
    review, _ = _review(build, db, EXIT_CODES, [
        "25. Termination",
        "25.1 The Customer may terminate for convenience by providing thirty (30) days' "
        "written notice.",
        "25.2 Leapswitch may terminate for convenience by providing ninety (90) days' written "
        "notice."])
    assert _observed(db, review, "NOTICE_TWO_WAYS") == []


# ------------------------------------------------------------------ missing and inapplicable
def test_a_clause_missing_in_its_own_family_is_missing_with_no_false_evidence(build, db):
    review, _ = _review(build, db, ["DATA-PURGE-MSA-001"], [
        "1. Services", "1.1 Leapswitch shall provide the Services described in the Order."])
    purge = _findings(db, review)["DATA-PURGE-MSA-001"]
    assert purge["classification"] == "MISSING" and purge["cited"] == []


INAPPLICABLE = [          # the shapes our own MSA has (5.2.2, 5.2.3, the bare heading 7)
    "5. Customer Onboarding",
    "5.2.2 All placed orders shall be subject to Know Your Customer (KYC) verification "
    "within twenty-four (24) hours.",
    "5.2.3 KYC Compliance for Resold Services: where the Customer resells the Services, it "
    "shall keep KYC records of its end customers.",
    "7. TERM AND TERMINATION",
    "7.1 This Agreement shall commence on the Service Commencement Date.",
]


def _uncertain(build, db, code):
    review, run = _review(build, db, [code], INAPPLICABLE)
    coverage = {c["code"]: c for c in run.applicability}
    return _findings(db, review)[code], coverage[code]["reason"]


def test_a_tos_retention_standard_is_never_missing_on_an_msa_from_a_verification_clause(
        build, db):
    """Step 1 (2026-10-21): "KYC verification" confirms the topic, not the TOS retention
    period; with no period stated the point is for Legal to confirm, never MISSING."""
    found, reason = _uncertain(build, db, "KYC-RETENTION-TOS-001")
    assert found["classification"] == "UNABLE_TO_EVALUATE"
    assert found["actual"].get("reason") == "APPLICABILITY_UNCERTAIN"
    assert "Legal to confirm it applies" in reason


def test_an_nda_termination_standard_is_never_missing_on_an_msa(build, db):
    """The bare heading "TERM AND TERMINATION" confirmed the NDA standard (score 5)."""
    found, _ = _uncertain(build, db, "TERM-NOTICE-NDA-001")
    assert found["classification"] == "UNABLE_TO_EVALUATE"
    assert found["actual"].get("reason") == "APPLICABILITY_UNCERTAIN"


def test_another_types_standard_whose_figure_is_stated_is_still_measured(build, db):
    review, _ = _review(build, db, ["TERM-NOTICE-NDA-001"], [
        "7. Term and Termination", "7.2 This Agreement may be terminated by either party on "
        "thirty (30) days' written notice."])
    assert _findings(db, review)["TERM-NOTICE-NDA-001"]["classification"] == "MATCH"


# ------------------------------------------------------------------ the real paper, locally
SOURCE = Path(os.environ.get("LEGALMIND_SOURCE_MATERIAL_DIR",
                             Path(__file__).resolve().parents[3] / "legal-docs"))


@pytest.mark.skipif(not (SOURCE / "MSA.pdf").is_file(),
                    reason="source material is gitignored (54.6); present on the server only")
def test_our_msa_data_purge_is_15_days_at_7_6_6_on_page_8(build, db):
    _ratified(build, "DATA-PURGE-MSA-001")
    contract = M.Contract(owner_id=build.owner.id, name="MSA", contract_type="MSA",
                          status=E.ContractStatus.ACTIVE)
    db.add(contract); db.flush()
    result = ingest_document(db, build.storage, contract_id=contract.id,
                             uploaded_by=build.owner.id,
                             data=(SOURCE / "MSA.pdf").read_bytes(), filename="MSA.pdf",
                             declared_mime="application/pdf")
    review = M.Review(contract_id=contract.id, document_version_id=result.document_version.id,
                      configuration_snapshot_id=build.snapshot.id,
                      status=E.ReviewStatus.DRAFT, created_by=build.owner.id)
    db.add(review); db.flush()
    run_analysis(db, review)
    purge = _findings(db, review)["DATA-PURGE-MSA-001"]
    assert purge["classification"] == "DEVIATION" and ("7.6.6", 8) in purge["cited"]


def test_an_other_typed_document_still_has_its_missing_cap(build, db):
    """OTHER is the catch-all (§31.13): an MSA position it lacks is still MISSING — the
    batch run of 2026-10-21 caught a bank SLA losing "no liability cap" to the new rule."""
    review, _ = _review(build, db, [CAP], [
        "5. Service Credits", "5.1 Service credits will not exceed the monthly fee."],
        kind="OTHER")
    assert _findings(db, review)[CAP]["classification"] == "MISSING"


# ------------------------------------------------------------------ Step 3: one case per finding
def _case(db, review, code, legal_position=True):
    angles = review_angles(db, review, legal_position=legal_position)
    return next(i["case"] for i in angles["legal"]["items"] if i["requirement_code"] == code)


def test_the_notice_case_states_contract_standard_gap_related_and_a_recommendation(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    case = _case(db, review, "CONVENIENCE-NOTICE-MSA-001")
    assert (case["contract"]["clause"], case["contract"]["page"]) == ("13.1", 2)
    assert "fifteen (15) days' written notice" in case["contract"]["quote"]
    assert case["gap"] == "15 days in the contract vs our 30 days"
    assert (case["standard"]["section"], case["standard"]["value"]) == ("13", "30 days")
    assert {r["clause"] for r in case["related"]} == {"13.2", "13.3"}
    assert case["recommendation"] in ("Request revised wording", "Seek legal review")


def test_a_one_sided_cap_case_is_not_negotiable_and_shows_no_rupees(build, db):
    review, _ = _cap_case(build, db, ONE_SIDED)
    case = _case(db, review, CAP)
    assert case["gap"] == "The cap binds only LeapSwitch; our standard is a mutual cap"
    assert case["recommendation"].startswith("Request revised wording — not negotiable")
    assert not (case["business"] or {}).get("money")


def test_an_uncertain_standard_case_recommends_legal_review(build, db):
    review, _ = _review(build, db, ["KYC-RETENTION-TOS-001"], INAPPLICABLE)
    case = _case(db, review, "KYC-RETENTION-TOS-001")
    assert case["gap"] == "May not apply to this document — Legal to confirm"
    assert case["recommendation"] == "Seek legal review"


def test_a_reader_without_positions_gets_the_case_without_our_value_or_source(build, db):
    review, _ = _review(build, db, EXIT_CODES, FIXED_TERM_EXIT, kind="OTHER")
    case = _case(db, review, "CONVENIENCE-NOTICE-MSA-001", legal_position=False)
    assert "value" not in case["standard"] and "source" not in case["standard"]
    assert case["gap"] == "15 days in the contract"


def test_no_case_ever_recommends_reject(build, db):
    review, _ = _review(build, db, [CAP, *EXIT_CODES, "DATA-PURGE-MSA-001"],
                        FIXED_TERM_EXIT, kind="MSA")
    for item in _angles(db, review)["legal"]["items"]:
        assert "reject" not in item["case"]["recommendation"].lower()


def test_an_unread_case_says_why_never_just_not_readable(build, db):
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability", "17.2 The total aggregate liability of each party shall "
        "not exceed the fees paid in the six (6) months preceding the claim for the specific "
        "services giving rise to the claim."])
    case = _case(db, review, CAP)
    assert case["gap"] == ("6 months in the contract, on a basis or unit our standard does "
                           "not use")


def test_a_clause_making_us_liable_for_data_loss_is_never_a_match(build, db):
    """The opposite of our exclusion (MSA3's suspension clause, 2026-10-21)."""
    review, _ = _review(build, db, ["DATA-LOSS-EXCLUSION-MSA-001"], [
        "6. Suspension", "6.4 LeapSwitch will be liable for any loss of data, content, or files "
        "due to suspension."])
    assert _findings(db, review)["DATA-LOSS-EXCLUSION-MSA-001"]["classification"] != "MATCH"


def test_a_contradiction_says_what_each_clause_states_and_where(build, db):
    """Owner, 2026-10-21: "how do they know why" — the reason names both statements."""
    review, _ = _review(build, db, [CAP], [
        "17. Limitation of Liability",
        "17.2 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the six (6) months preceding the claim.",
        "17.7 The total aggregate liability of each party shall not exceed the total fees "
        "paid by the Customer in the twelve (12) months preceding the claim."])
    gap = _case(db, review, CAP)["gap"]
    assert "6 months and 12 months" in gap and "17.2" in gap and "17.7" in gap
    assert gap.endswith("which governs is for a person to decide")


def test_an_nda_position_expected_beside_confidentiality_is_not_missing_on_an_msa(build, db):
    """Owner, 2026-10-21: a lawyer does not raise "Residuals missing" on an MSA; the
    coverage record says why, and Legal may add it."""
    _, run = _review(build, db, ["CONF-SURVIVAL-MSA-001", "RESIDUALS-NDA-001"], [
        "15. Confidentiality",
        "15.4 The obligations of confidentiality shall survive for three (3) years after "
        "the termination of this Agreement."])
    coverage = {c["code"]: c for c in run.applicability}
    assert coverage["RESIDUALS-NDA-001"]["outcome"] == "NOT_APPLICABLE"
    assert "Legal may add it" in coverage["RESIDUALS-NDA-001"]["reason"]


DATA_LOSS = "DATA-LOSS-EXCLUSION-MSA-001"


def test_a_data_loss_exclusion_limited_to_customer_causes_is_narrower(build, db):
    """G-F (owner, 2026-10-21): our MSA §9.9's shape — the position also covers our own
    operational error (Constitution L1.14 §9)."""
    review, _ = _review(build, db, [DATA_LOSS], [
        "9. Liability", "9.9 Leapswitch shall not be liable for any direct or indirect "
        "damages, including without limitation loss of data, arising from customer-supplied "
        "misinformation, misuse, third-party systems, or incompatible hardware/software "
        "configurations."])
    found = _findings(db, review)[DATA_LOSS]
    assert (found["classification"], found["actual"].get("presence")) == (
        "DEVIATION", "TOO_NARROW")
    assert _case(db, review, DATA_LOSS)["gap"].startswith("Narrower than our standard")


def test_a_full_data_loss_exclusion_still_matches(build, db):
    review, _ = _review(build, db, [DATA_LOSS], [
        "9. Liability", "9.9 Leapswitch shall not be liable for any loss, deletion, corruption "
        "or unavailability of customer data, howsoever caused, including its own operational "
        "error."])
    assert _findings(db, review)[DATA_LOSS]["classification"] == "MATCH"


def test_an_exclusion_arising_out_of_the_agreement_is_not_narrowed(build, db):
    review, _ = _review(build, db, [DATA_LOSS], [
        "13. Limitation", "13.1 In no event will either Party be liable for any loss of data "
        "arising out of or relating to this Agreement."])
    assert _findings(db, review)[DATA_LOSS]["actual"].get("presence") != "TOO_NARROW"

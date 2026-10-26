"""Word's automatic list numbering is the document's own clause numbering (AM-185, owner
2026-10-26). A Word MSA read 0 of 219 numbered clauses while its PDF read 213, so nothing
could be cited and clauses ran into each other. Synthetic files only (54.6): each case is
a way real Word documents number their clauses, with the number Word itself displays."""
from __future__ import annotations

import io

import docx
from docx.oxml import parse_xml

from legalmind.ingestion.parsing import parse

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _lvl(ilvl: int, fmt: str, text: str, start: int = 1) -> str:
    return (f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/>'
            f'<w:lvlText w:val="{text}"/></w:lvl>')


def _document(build) -> bytes:
    """A document whose numbering part holds: abstract 10 = 1. / 1.1 / (a); abstract 11
    = a bullet; abstract 12 = "%1" for numbered headings; num 1 and num 2 share abstract
    10 (num 2 restarts at 1), num 3 = bullets, num 4 = headings."""
    document = docx.Document()
    numbering = document.part.numbering_part.element
    for xml in (
        f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="10">{_lvl(0, "decimal", "%1.")}'
        f'{_lvl(1, "decimal", "%1.%2")}{_lvl(2, "lowerLetter", "(%3)")}</w:abstractNum>',
        f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="11">{_lvl(0, "bullet", "•")}'
        '</w:abstractNum>',
        f'<w:abstractNum xmlns:w="{W}" w:abstractNumId="12">{_lvl(0, "decimal", "%1")}'
        '</w:abstractNum>',
    ):
        numbering.insert(0, parse_xml(xml))
    for xml in (
        f'<w:num xmlns:w="{W}" w:numId="1"><w:abstractNumId w:val="10"/></w:num>',
        f'<w:num xmlns:w="{W}" w:numId="2"><w:abstractNumId w:val="10"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>',
        f'<w:num xmlns:w="{W}" w:numId="3"><w:abstractNumId w:val="11"/></w:num>',
        f'<w:num xmlns:w="{W}" w:numId="4"><w:abstractNumId w:val="12"/></w:num>',
    ):
        numbering.append(parse_xml(xml))

    def item(text: str, num: int, level: int = 0, where=None):
        p = (where or document).add_paragraph(text)
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr xmlns:w="{W}"><w:ilvl w:val="{level}"/><w:numId w:val="{num}"/>'
            '</w:numPr>'))
        return p

    build(document, item)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _numbers(data: bytes) -> list[tuple[str | None, str]]:
    return [(s.section_number, s.content.split("\n")[0][:30])
            for s in parse(data, DOCX).segments]


def test_word_numbers_a_multi_level_contract_as_it_displays_it():
    def build(d, item):
        item("Definitions", 1)
        item("Affiliate means any person controlling a party.", 1, 1)
        item("Business Day means a day banks are open.", 1, 1)
        item("including its successors", 1, 2)              # (a): part of 1.2
        item("Payment", 1)
        item("Invoices are payable within 21 days.", 1, 1)
    assert [n for n, _ in _numbers(_document(build))] == ["1", "1.1", "1.2", None, "2",
                                                          "2.1"]


def test_a_number_the_text_states_is_kept_over_the_list_number():
    def build(d, item):
        item("7.3 Liability is capped at fees paid in 12 months.", 1)
    assert _numbers(_document(build))[0][0] == "7.3"


def test_an_annexure_list_that_restarts_starts_again_at_one():
    def build(d, item):
        item("Term", 1)
        item("Fees", 1)
        d.add_paragraph("ANNEXURE A")
        item("Rack space", 2)                               # num 2: startOverride 1
        item("Power", 2)
    assert [n for n, _ in _numbers(_document(build))] == ["1", "2", None, "1", "2"]


def test_bullets_carry_no_clause_number():
    def build(d, item):
        item("Confidentiality", 1)
        item("No liability for indirect damages", 3)
        item("Liability capped at fees paid", 3)
    assert [n for n, _ in _numbers(_document(build))] == ["1", None, None]


def test_a_heading_style_that_carries_numbering_is_numbered():
    def build(d, item):
        style = d.styles["Heading 1"]
        style.element.get_or_add_pPr().append(parse_xml(
            f'<w:numPr xmlns:w="{W}"><w:ilvl w:val="0"/><w:numId w:val="4"/></w:numPr>'))
        d.add_paragraph("Scope", style="Heading 1")
        d.add_paragraph("The Vendor provides colocation.")
        d.add_paragraph("Fees", style="Heading 1")
    assert [n for n, _ in _numbers(_document(build))] == ["1", None, "2"]


def test_a_numbered_paragraph_inside_a_table_still_counts():
    """Word counts it, so the clause after the table is numbered as Word shows it."""
    def build(d, item):
        item("Scope", 1)
        cell = d.add_table(rows=1, cols=1).cell(0, 0)
        item("Services", 1, where=cell)
        item("Fees", 1)
    numbers = [n for n, _ in _numbers(_document(build))]
    assert numbers[0] == "1" and numbers[-1] == "3"


def test_numbering_switched_off_is_no_number():
    def build(d, item):
        item("Plain paragraph", 0)                           # numId 0: no numbering
    assert _numbers(_document(build))[0][0] is None


# Break tests: numbering Word files really carry that must never invent a number or
# fail the upload.

def test_a_roman_or_letter_level_is_never_a_clause_number():
    def build(d, item):
        item("Scope", 1)
        item("rack space", 1, 2)                             # (a)
    assert [n for n, _ in _numbers(_document(build))] == ["1", None]


def test_a_reference_to_a_list_that_does_not_exist_is_no_number():
    def build(d, item):
        item("Scope", 99)
    assert _numbers(_document(build))[0][0] is None


def test_a_level_the_list_does_not_define_is_no_number():
    def build(d, item):
        item("Scope", 1, 7)
    assert _numbers(_document(build))[0][0] is None


def test_malformed_numbering_reads_the_document_without_numbers():
    def build(d, item):
        lvl = d.part.numbering_part.element.find(f"{{{W}}}abstractNum/{{{W}}}lvl")
        lvl.find(f"{{{W}}}start").set(f"{{{W}}}val", "x")
        item("Scope", 1)
        d.add_paragraph("The Vendor provides colocation.")
    assert [n for n, _ in _numbers(_document(build))] == [None, None]


def test_a_list_numbered_clause_body_is_never_its_own_heading():
    """The whole paragraph as the title made every heading term match the body: a price
    term matched the confidentiality clause (replay, 2026-10-26)."""
    body = ("The Parties shall endeavor to protect Confidential Information, and any fees "
            "or price disclosed under this Agreement is Confidential Information of the "
            "party that discloses it, whatever form it takes and however it is marked, for "
            "as long as this Agreement remains in force.")
    def build(d, item):
        item("CONFIDENTIALITY", 1)
        item(body, 1, 1)
    first, second = parse(_document(build), DOCX).segments
    assert (first.section_number, first.section_title) == ("1", "CONFIDENTIALITY")
    assert second.section_number == "1.1" and second.section_title is None


def test_a_short_list_numbered_clause_is_not_a_heading_either():
    """17.7's one-line cap read as a heading made "fees" a heading term (replay)."""
    def build(d, item):
        item("LIABILITY", 1)
        item("Total liability shall not exceed the fees paid in the 12 months before.", 1, 1)
    first, second = parse(_document(build), DOCX).segments
    assert first.section_title == "LIABILITY" and second.section_title is None
